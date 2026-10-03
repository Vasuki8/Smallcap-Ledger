"""Read-only BOI TER recovery from real shared-string, styled XLSX packages."""
import hashlib
import io
import unittest
import zipfile
from datetime import date, datetime
from unittest.mock import patch
from xml.etree import ElementTree as ET

import openpyxl

from tracker import midcap_ter_first_party as first_party
from tracker.midcap_ter_readiness import reconcile


FAMILY = "BANK OF INDIA MID CAP FUND"
PUBLISHED_FAMILY = "Bank of India Mid Cap Fund"
AMC = "Bank of India Mutual Fund"
NSDL = "BOIA/O/E/MIF/25/06/0023"
SOURCE = (
    "https://www.boimf.in/docs/default-source/investorcorner/total-expense-ratio/"
    "expense_ratio_01092026_to_30092026.xls?sfvrsn=1b375908_8"
)
TODAY = date(2026, 10, 3)
HEADER = [
    "NSDL Scheme Code", "Scheme Name", "TER Date\n(DD/MM/\nYYYY)",
    "Regular Plan - Base Expense Ratio (BER) (%)", "Regular Plan - Brokerage cost (%)",
    "Regular Plan - Transaction Cost incurred for the purpose of execution of trade (%)",
    "Regular Plan - Statutory Levies (including GST) (%)", "Regular Plan - Total TER (%)",
    "Direct Plan - Base Expense Ratio (BER) (%)", "Direct Plan - Brokerage cost (%)",
    "Direct Plan - Transaction Cost incurred for the purpose of execution of trade (%)",
    "Direct Plan - Statutory Levies (including GST) (%)", "Direct Plan - Total TER (%)",
]


def scheme_row(day=date(2026, 9, 29), **changes):
    row = dict(zip("ABCDEFGHIJKLM", [
        NSDL, PUBLISHED_FAMILY, day, 2.02, 0.06, 0.0, 0.47, 2.55,
        0.98, 0.06, 0.0, 0.28, 1.32,
    ]))
    return {**row, **changes}


def workbook(*rows, sheet_names=("TER_UPLOAD_FORMAT",), header_changes=None, formats=None, raw_values=None):
    """Keep genuine Excel styles while replaying the source's shared strings/A1."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    book.properties.created = book.properties.modified = datetime(2026, 10, 3)
    headers = list(HEADER)
    for column, value in (header_changes or {}).items():
        headers[ord(column) - ord("A")] = value
    for name in sheet_names:
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for values in rows:
            sheet.append([values.get(column) for column in "ABCDEFGHIJKLM"])
        for row in sheet.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=1, max_col=13):
            row[2].number_format = "dd/MM/yyyy"
            for cell in row[3:]:
                cell.number_format = "#,##0.00"
        for coordinate, number_format in (formats or {}).items():
            sheet[coordinate].number_format = number_format
    saved = io.BytesIO()
    book.save(saved)
    book.close()
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    string_values = []
    with zipfile.ZipFile(io.BytesIO(saved.getvalue())) as package:
        files = {name: package.read(name) for name in package.namelist()}
    for name in list(files):
        if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
            continue
        root = ET.fromstring(files[name])
        root.find(f"{{{ns}}}dimension").set("ref", "A1")
        for cell in root.findall(f".//{{{ns}}}c"):
            if cell.get("t") == "inlineStr":
                inline = cell.find(f"{{{ns}}}is")
                if inline is None:
                    continue
                value = "".join(node.text or "" for node in inline.iter(f"{{{ns}}}t"))
                if value not in string_values:
                    string_values.append(value)
                cell.remove(inline)
                cell.set("t", "s")
                ET.SubElement(cell, f"{{{ns}}}v").text = str(string_values.index(value))
            if cell.get("r") in (raw_values or {}):
                cell.set("t", "n")
                value_node = cell.find(f"{{{ns}}}v")
                if value_node is None:
                    value_node = ET.SubElement(cell, f"{{{ns}}}v")
                value_node.text = raw_values[cell.get("r")]
        files[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    strings = ET.Element(f"{{{ns}}}sst", count=str(len(string_values)), uniqueCount=str(len(string_values)))
    for value in string_values:
        ET.SubElement(ET.SubElement(strings, f"{{{ns}}}si"), f"{{{ns}}}t").text = value
    files["xl/sharedStrings.xml"] = ET.tostring(strings, encoding="utf-8", xml_declaration=True)
    content_types = ET.fromstring(files["[Content_Types].xml"])
    ET.SubElement(content_types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
        "PartName": "/xl/sharedStrings.xml",
        "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
    })
    files["[Content_Types].xml"] = ET.tostring(content_types, encoding="utf-8", xml_declaration=True)
    relationships = ET.fromstring(files["xl/_rels/workbook.xml.rels"])
    ET.SubElement(relationships, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", {
        "Id": "rIdSharedStrings", "Target": "sharedStrings.xml",
        "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
    })
    files["xl/_rels/workbook.xml.rels"] = ET.tostring(relationships, encoding="utf-8", xml_declaration=True)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as package:
        for name, content in files.items():
            info = zipfile.ZipInfo(name, date_time=(2026, 10, 3, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            package.writestr(info, content)
    return output.getvalue()


class BoiMidcapTerTests(unittest.TestCase):
    def collect_source(self, content, *, today=TODAY, staged=True):
        requests = []

        def fetch(source, **kwargs):
            requests.append(source)
            self.assertEqual(source, SOURCE)
            self.assertIs(kwargs.get("archive"), False)
            self.assertEqual(kwargs.get("max_bytes"), 5 * 1024 * 1024)
            if isinstance(content, Exception):
                raise content
            return content, "0" * 64, "application/vnd.ms-excel"

        collectors = {family: collector for family, collector in first_party.COLLECTORS.items() if family == FAMILY}
        with patch.object(first_party, "COLLECTORS", collectors), \
             patch.object(first_party.db, "rows", return_value=[{"family": FAMILY}] if staged else []), \
             patch.object(first_party.providers, "fetch", side_effect=fetch), \
             patch.object(first_party.db, "connect", side_effect=AssertionError("Database connections forbidden")) as connection:
            report = first_party.collect(today=today)
        connection.assert_not_called()
        self.assertEqual(report["production_writes"], 0)
        self.assertFalse(report["public_export_enabled"])
        return report, requests

    def assert_provenance(self, record, content):
        self.assertEqual(record.get("source"), SOURCE)
        self.assertEqual(record.get("sha256"), hashlib.sha256(content).hexdigest())
        self.assertEqual(record.get("bytes"), len(content))

    def assert_recovered(self, report, content, *, as_of="2026-09-29", direct=1.32, regular=2.55):
        self.assertEqual(report["recovered"], 1)
        self.assertEqual(report["errors"], [])
        result = report["results"][0]
        self.assertEqual(result["family"], FAMILY)
        self.assertEqual(result["amc"], AMC)
        self.assertEqual(result["status"], "recovered")
        self.assertEqual(result["as_of"], as_of)
        self.assertEqual(result["direct_ter"], direct)
        self.assertEqual(result["regular_ter"], regular)
        self.assertEqual(result["identity"], {"scheme_name": PUBLISHED_FAMILY, "nsdl_scheme_code": NSDL})
        self.assertEqual(result["discovery_channel"], "reviewed_monthly_workbook")
        self.assert_provenance(result, content)
        return result

    def assert_failed(self, report, *, content=None, message=None):
        self.assertEqual(report["results"], [])
        self.assertEqual(report["failed"], 1)
        error = report["errors"][0]
        self.assertEqual(error["family"], FAMILY)
        self.assertTrue(error.get("error"))
        if message:
            self.assertRegex(error["error"].casefold(), message)
        if content is not None:
            self.assert_provenance(error, content)
        return error

    def test_registered_collector_recovers_shared_strings_styled_dates_beyond_a1_dimension(self):
        content = workbook(scheme_row(date(2026, 9, 28)), scheme_row())
        report, requests = self.collect_source(content)
        result = self.assert_recovered(report, content)
        self.assertEqual(requests, [SOURCE])
        self.assertEqual(result["plans"]["Direct"]["transaction_cost"], 0.0)
        self.assertEqual(result["plans"]["Regular"]["transaction_cost"], 0.0)

    def test_source_is_eligible_in_its_current_calendar_month(self):
        content = workbook(scheme_row())
        report, requests = self.collect_source(content, today=date(2026, 9, 29))
        self.assert_recovered(report, content)
        self.assertEqual(requests, [SOURCE])

    def test_expired_source_period_fails_before_fetch_without_fabricated_provenance(self):
        report, requests = self.collect_source(workbook(scheme_row()), today=date(2026, 11, 1))
        error = self.assert_failed(report, message="period|month|eligible|reviewed")
        self.assertEqual(requests, [])
        self.assertEqual(error.get("source"), SOURCE)
        self.assertTrue({"sha256", "bytes"}.isdisjoint(error))

    def test_collector_requires_staged_family_before_any_fetch(self):
        report, requests = self.collect_source(workbook(scheme_row()), staged=False)
        error = self.assert_failed(report, message="staged")
        self.assertEqual(requests, [])
        self.assertTrue({"sha256", "bytes"}.isdisjoint(error))

    def test_exact_sheet_and_required_headers_fail_closed(self):
        cases = [
            ("wrong_sheet", workbook(scheme_row(), sheet_names=("Other",)), "sheet"),
            ("extra_sheet", workbook(scheme_row(), sheet_names=("TER_UPLOAD_FORMAT", "Other")), "sheet"),
            ("nsdl_header_changed", workbook(scheme_row(), header_changes={"A": "Scheme Code"}), "header|column|schema"),
            ("direct_ter_header_changed", workbook(scheme_row(), header_changes={"M": "Direct Plan - BER (%)"}), "header|column|schema"),
        ]
        for label, content, message in cases:
            with self.subTest(reason=label):
                report, _ = self.collect_source(content)
                self.assert_failed(report, content=content, message=message)

    def test_family_and_nsdl_contradictions_cannot_be_filtered_away(self):
        valid = scheme_row(date(2026, 9, 28))
        cases = [
            ("wrong_nsdl_for_family", scheme_row(A="BOIA/O/E/SCF/18/11/0014")),
            ("missing_nsdl_for_family", scheme_row(A=None)),
            ("empty_nsdl_for_family", scheme_row(A="")),
            ("expected_nsdl_wrong_family", scheme_row(B="Bank of India Small Cap Fund")),
        ]
        for label, invalid in cases:
            with self.subTest(reason=label):
                content = workbook(valid, invalid)
                report, _ = self.collect_source(content)
                self.assert_failed(report, content=content, message="identity|nsdl|family|scheme")

    def test_matching_identity_date_errors_do_not_rescue_an_earlier_valid_row(self):
        valid = scheme_row(date(2026, 9, 27))
        cases = [
            ("missing", workbook(valid, scheme_row(C=None)), TODAY),
            ("text_date", workbook(valid, scheme_row(C="29/09/2026")), TODAY),
            ("unstyled_serial", workbook(valid, scheme_row(), formats={"C3": "General"}), TODAY),
            ("outside_source_period", workbook(valid, scheme_row(date(2026, 10, 1))), TODAY),
            ("future", workbook(valid, scheme_row()), date(2026, 9, 28)),
        ]
        for label, content, today in cases:
            with self.subTest(reason=label):
                report, _ = self.collect_source(content, today=today)
                self.assert_failed(report, content=content, message="date|period|future")

    def test_duplicate_latest_exact_rows_are_rejected(self):
        content = workbook(scheme_row(date(2026, 9, 28)), scheme_row(), scheme_row())
        report, _ = self.collect_source(content)
        self.assert_failed(report, content=content, message="duplicate|unique")

    def test_invalid_latest_totals_cannot_promote_an_earlier_valid_row(self):
        cases = [
            ("components", scheme_row(M=1.29), "reconcile|component"),
            ("ter_below_ber", scheme_row(I=1.4), "below|ber"),
        ]
        for label, latest, message in cases:
            with self.subTest(reason=label):
                content = workbook(scheme_row(date(2026, 9, 28)), latest)
                report, _ = self.collect_source(content)
                self.assert_failed(report, content=content, message=message)

    def test_missing_placeholder_and_out_of_range_numeric_components_are_rejected(self):
        values = [None, "", "-", "NA", "N/A", "not numeric", -0.01, 5.01]
        for invalid in values:
            with self.subTest(value=invalid):
                content = workbook(scheme_row(F=invalid))
                report, _ = self.collect_source(content)
                self.assert_failed(report, content=content, message="number|numeric|invalid|missing|range")

    def test_booleans_are_not_accepted_as_numeric_zero_or_one(self):
        for value, total in ((False, 2.55), (True, 3.55)):
            with self.subTest(value=value):
                content = workbook(scheme_row(F=value, H=total))
                report, _ = self.collect_source(content)
                self.assert_failed(report, content=content, message="number|numeric|bool|invalid")

    def test_nonfinite_numeric_xml_values_are_rejected_with_acquired_provenance(self):
        for raw in ("1E309", "-1E309", "NaN"):
            with self.subTest(raw=raw):
                content = workbook(scheme_row(), raw_values={"F2": raw})
                report, _ = self.collect_source(content)
                self.assert_failed(report, content=content, message="number|numeric|invalid|finite|range")

    def test_percentage_formatted_fractions_cannot_silently_change_units(self):
        fractional = scheme_row()
        for column in "DEFGHIJKLM":
            fractional[column] /= 100
        content = workbook(fractional, formats={f"{column}2": "0.00%" for column in "DEFGHIJKLM"})
        report, _ = self.collect_source(content)
        self.assert_failed(report, content=content, message="percentage|percent|format|unit")

    def test_another_valid_snapshot_uses_actual_row_date_values_and_hash(self):
        content = workbook(scheme_row(date(2026, 9, 30), L=0.29, M=1.33))
        report, _ = self.collect_source(content)
        self.assert_recovered(report, content, as_of="2026-09-30", direct=1.33)

    def test_transport_failure_has_source_without_fabricated_hash_or_size(self):
        report, requests = self.collect_source(RuntimeError("BOI workbook unavailable"))
        error = self.assert_failed(report, message="unavailable")
        self.assertEqual(requests, [SOURCE])
        self.assertEqual(error.get("source"), SOURCE)
        self.assertTrue({"sha256", "bytes"}.isdisjoint(error))

    def test_invalid_package_still_keeps_the_actual_acquired_byte_provenance(self):
        content = b"not an XLSX package"
        report, _ = self.collect_source(content)
        self.assert_failed(report, content=content)

    def test_rejected_source_evidence_remains_excluded_from_readiness(self):
        content = workbook(scheme_row(M=1.29))
        report, _ = self.collect_source(content)
        self.assert_failed(report, content=content)
        readiness = reconcile({"families": [{
            "family": FAMILY, "amc": AMC, "direct_ter_available": False, "ter": None,
        }]}, report)
        self.assertEqual(readiness["counts"]["direct_ter"], 0)
        self.assertEqual(readiness["counts"]["first_party_amc"], 0)
        self.assertEqual(readiness["remaining_families"], [FAMILY])
        self.assertFalse(readiness["families"][0]["direct_ter_available"])
        self.assertIsNone(readiness["families"][0]["direct_ter"])
        self.assertFalse(readiness["public_export_enabled"])


if __name__ == "__main__":
    unittest.main()
