import io
import json
import unittest
import zipfile
from xml.sax.saxutils import escape

import openpyxl
from datetime import date
from unittest.mock import patch

from tracker import amc_expenses


class CanaraExpenseTests(unittest.TestCase):
    def rows(self, day="2026-09-24"):
        return [
            {
                "id": "1",
                "sch_code": "SC",
                "scheme_name": "Canara Robeco Small Cap Fund",
                "date": day,
                "plan_type": "Regular Plan",
                "base_ter": "1.46",
                "additional_expense_6a_b": "0.05",
                "additional_expense_6a_c": "0",
                "gst": "0.33",
                "total_ter": "1.84",
            },
            {
                "id": "2",
                "sch_code": "SC",
                "scheme_name": "Canara Robeco Small Cap Fund",
                "date": day,
                "plan_type": "Direct Plan",
                "base_ter": "0.46",
                "additional_expense_6a_b": "0.05",
                "additional_expense_6a_c": "0",
                "gst": "0.17",
                "total_ter": "0.68",
            },
        ]

    def test_exact_latest_plan_pair_retains_published_ber_and_ter(self):
        records = self.rows("2026-09-23") + self.rows("2026-09-24")
        day, plans = amc_expenses.parse_canara_records(records, date(2026, 9, 24))
        self.assertEqual(day, "2026-09-24")
        self.assertEqual(
            plans,
            {
                "Regular": {"base_expense_ratio": 1.46, "ter": 1.84},
                "Direct": {"base_expense_ratio": 0.46, "ter": 0.68},
            },
        )

    def test_future_rows_are_ignored_and_latest_complete_day_is_used(self):
        records = self.rows("2026-09-23") + self.rows("2026-09-25")
        day, plans = amc_expenses.parse_canara_records(records, date(2026, 9, 24))
        self.assertEqual(day, "2026-09-23")
        self.assertEqual(plans["Direct"]["ter"], 0.68)

    def test_wrong_scheme_or_code_cannot_be_promoted(self):
        rows = self.rows()
        rows[0]["scheme_name"] = "Canara Robeco Mid Cap Fund"
        with self.assertRaisesRegex(ValueError, "no complete dated"):
            amc_expenses.parse_canara_records(rows, date(2026, 9, 24))

        rows = self.rows()
        rows[1]["sch_code"] = "OTHER"
        with self.assertRaisesRegex(ValueError, "no complete dated"):
            amc_expenses.parse_canara_records(rows, date(2026, 9, 24))

    def test_duplicate_plan_or_total_below_ber_is_rejected(self):
        rows = self.rows()
        rows.append(dict(rows[1], id="3"))
        with self.assertRaisesRegex(ValueError, "duplicate plan rows"):
            amc_expenses.parse_canara_records(rows, date(2026, 9, 24))

        rows = self.rows()
        rows[1]["total_ter"] = "0.40"
        with self.assertRaisesRegex(ValueError, "below BER"):
            amc_expenses.parse_canara_records(rows, date(2026, 9, 24))

    @patch("tracker.amc_expenses.db.metric")
    @patch("tracker.amc_expenses.db.one", return_value={"code": "CANARA"})
    @patch("tracker.amc_expenses.fetch")
    def test_live_collector_shape_stores_four_exact_metrics(self, mock_fetch, _mock_one, mock_metric):
        body = json.dumps(self.rows()).encode()
        mock_fetch.return_value = (body, "hash123", "application/json")
        result = amc_expenses.canara(today=date(2026, 9, 24), lookback_days=3)

        self.assertIn("2026-09-24", result)
        url = mock_fetch.call_args.args[0]
        self.assertIn("from_date=2026-09-22", url)
        self.assertIn("to_date=2026-09-24", url)
        self.assertEqual(mock_metric.call_count, 4)

        calls = {
            (call.args[1], call.args[2]): call.args[4]
            for call in mock_metric.call_args_list
        }
        self.assertEqual(calls[("Regular", "base_expense_ratio")], 1.46)
        self.assertEqual(calls[("Regular", "ter")], 1.84)
        self.assertEqual(calls[("Direct", "base_expense_ratio")], 0.46)
        self.assertEqual(calls[("Direct", "ter")], 0.68)
        self.assertTrue(all(call.args[3] == "2026-09-24" for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[7] == "hash123" for call in mock_metric.call_args_list))


class GrowwExpenseTests(unittest.TestCase):
    def notice_text(
        self,
        *,
        as_of="September 23, 2026",
        signed="September 24, 2026",
        effective="September 30, 2026",
        row="Groww Smallcap Fund 0.42 1.94 0.49 1.94 (No change)",
        heading="Scheme(s) Name Current BER* Revised BER** Direct (%) Regular (%) Direct (%) Regular (%)",
        notice=26,
    ):
        return f"""
        Notice no. {notice}/2026 – 2027
        Change in Base Expense Ratio (‘BER’) of the scheme(s) of Groww Mutual Fund:
        NOTICE is hereby given that the BER shall be as follows with effect from {effective} (‘Effective date’):
        * As on {as_of}.
        ** Or such lower BER as may be applicable due to AUM on the Effective date.
        Authorised Signatory Date: {signed}
        {heading}
        {row}
        """

    def test_groww_current_ber_is_stored_and_future_revised_ber_is_not_promoted(self):
        day, plans, notice, signed, effective = amc_expenses.parse_groww_ber_text(
            self.notice_text(), date(2026, 9, 25)
        )
        self.assertEqual(day, "2026-09-23")
        self.assertEqual(plans, {"Direct": 0.42, "Regular": 1.94})
        self.assertEqual(notice, 26)
        self.assertEqual(signed, "2026-09-24")
        self.assertEqual(effective, "2026-09-30")
        self.assertNotIn(0.49, plans.values())

    def test_groww_na_regular_is_preserved_as_missing_not_inferred(self):
        day, plans, notice, _, _ = amc_expenses.parse_groww_ber_text(
            self.notice_text(
                as_of="August 31, 2026",
                signed="September 01, 2026",
                effective="September 05, 2026",
                row="Groww Smallcap Fund 0.45 NA 0.51 NA",
                notice=24,
            ),
            date(2026, 9, 25),
        )
        self.assertEqual(day, "2026-08-31")
        self.assertEqual(plans, {"Direct": 0.45})
        self.assertEqual(notice, 24)

    def test_groww_notice_identity_heading_and_dates_are_strict(self):
        with self.assertRaisesRegex(ValueError, "table heading changed"):
            amc_expenses.parse_groww_ber_text(
                self.notice_text(heading="Scheme(s) Name Current TER Revised TER"),
                date(2026, 9, 25),
            )
        with self.assertRaisesRegex(ValueError, "future observation"):
            amc_expenses.parse_groww_ber_text(
                self.notice_text(as_of="September 26, 2026"),
                date(2026, 9, 25),
            )
        with self.assertRaisesRegex(ValueError, "one exact Smallcap row"):
            amc_expenses.parse_groww_ber_text(
                self.notice_text(row="Groww Midcap Fund 0.42 1.94 0.49 1.94"),
                date(2026, 9, 25),
            )

    def test_groww_page_discovery_keeps_only_current_year_registered_ber_notices(self):
        html = b"""
        <a href="https://assets-netstorage.growwmf.in/compliance_docs/Downloads/Expense%20Ratio/Notice%20-%20Change%20in%20TER/2026%20-%202027/26.%20Notice%20-%20Change%20in%20BER.pdf">26. Notice - Change in BER.pdf</a>
        <a href="https://assets-netstorage.growwmf.in/compliance_docs/Downloads/Expense%20Ratio/Notice%20-%20Change%20in%20TER/2026%20-%202027/24.%20Notice%20-%20Change%20in%20BER.pdf">24. Notice - Change in BER.pdf</a>
        <a href="https://example.com/26.%20Notice%20-%20Change%20in%20BER.pdf">26. Notice - Change in BER.pdf</a>
        <a href="https://assets-netstorage.growwmf.in/compliance_docs/Downloads/Expense%20Ratio/Notice%20-%20Change%20in%20TER/2025%20-%202026/61.%20Notice%20-%20Change%20in%20BER.pdf">61. Notice - Change in BER.pdf</a>
        """
        links = amc_expenses._groww_ber_links(html, date(2026, 9, 25))
        self.assertEqual([x[0] for x in links], [26, 24])
        self.assertTrue(all("assets-netstorage.growwmf.in" in x[1] for x in links))

    @patch("tracker.amc_expenses.db.metric")
    @patch("tracker.amc_expenses.db.one", return_value={"code": "GROWW"})
    @patch("tracker.amc_expenses._groww_disclosure")
    def test_groww_collector_stores_only_current_ber_with_source_hash(
        self, mock_disclosure, _mock_one, mock_metric
    ):
        source = (
            "https://assets-netstorage.growwmf.in/compliance_docs/Downloads/"
            "Expense%20Ratio/Notice%20-%20Change%20in%20TER/2026%20-%202027/"
            "26.%20Notice%20-%20Change%20in%20BER.pdf"
        )
        mock_disclosure.return_value = (
            source,
            "2026-09-23",
            {"Direct": 0.42, "Regular": 1.94},
            "growwhash",
            "2026-09-24",
            "2026-09-30",
        )
        result = amc_expenses.groww(today=date(2026, 9, 25))
        self.assertIn("2026-09-23", result)
        self.assertIn("not promoted", result)
        self.assertEqual(mock_metric.call_count, 2)
        calls = {(call.args[1], call.args[2]): call.args[4] for call in mock_metric.call_args_list}
        self.assertEqual(calls[("Direct", "base_expense_ratio")], 0.42)
        self.assertEqual(calls[("Regular", "base_expense_ratio")], 1.94)
        self.assertTrue(all(call.args[3] == "2026-09-23" for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[6] == source for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[7] == "growwhash" for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[2] != "ter" for call in mock_metric.call_args_list))


class HsbcExpenseTests(unittest.TestCase):
    def workbook(self, rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TER"
        ws.append([None] * 14)
        ws.append(["Total Expense Ratio (TER) for HSBC Mutual Fund"] + [None] * 13)
        ws.append(list(amc_expenses._HSBC_HEADER))
        for row in rows:
            ws.append(row)
        out = io.BytesIO()
        wb.save(out)
        wb.close()
        return out.getvalue()

    def row(self, day="2026-09-23", *, scheme="HSBC Small Cap Fund", code="HEMIDF",
            nsdl="LTMF/O/E/SCF/14/02/0023", regular_ter=1.78, direct_ter=0.77):
        return [
            code, nsdl, scheme, day,
            1.42, 0.03, 0.00, 0.33, regular_ter,
            0.56, 0.03, 0.00, 0.18, direct_ter,
        ]

    def test_latest_exact_hsbc_row_retains_published_ber_and_total_ter(self):
        book = self.workbook([self.row("2026-09-22"), self.row("2026-09-23")])
        day, plans = amc_expenses.parse_hsbc_workbook(book, date(2026, 9, 24))
        self.assertEqual(day, "2026-09-23")
        self.assertEqual(plans["Regular"]["base_expense_ratio"], 1.42)
        self.assertEqual(plans["Regular"]["ter"], 1.78)
        self.assertEqual(plans["Direct"]["base_expense_ratio"], 0.56)
        self.assertEqual(plans["Direct"]["ter"], 0.77)

    def test_hsbc_future_row_is_ignored(self):
        book = self.workbook([self.row("2026-09-23"), self.row("2026-09-25")])
        day, plans = amc_expenses.parse_hsbc_workbook(book, date(2026, 9, 24))
        self.assertEqual(day, "2026-09-23")
        self.assertEqual(plans["Direct"]["ter"], 0.77)

    def test_hsbc_identity_is_exact(self):
        for kwargs in (
            {"scheme": "HSBC Midcap Fund"},
            {"code": "OTHER"},
            {"nsdl": "LTMF/O/E/OTHER"},
        ):
            book = self.workbook([self.row(**kwargs)])
            with self.assertRaisesRegex(ValueError, "no dated Small Cap rows"):
                amc_expenses.parse_hsbc_workbook(book, date(2026, 9, 24))

    def test_hsbc_duplicate_or_unreconciled_total_is_rejected(self):
        row = self.row()
        book = self.workbook([row, list(row)])
        with self.assertRaisesRegex(ValueError, "duplicate Small Cap rows"):
            amc_expenses.parse_hsbc_workbook(book, date(2026, 9, 24))

        book = self.workbook([self.row(direct_ter=0.80)])
        with self.assertRaisesRegex(ValueError, "components do not reconcile"):
            amc_expenses.parse_hsbc_workbook(book, date(2026, 9, 24))

    def test_hsbc_workbook_header_is_strict(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TER"
        ws.append([None] * 14)
        ws.append(["Total Expense Ratio (TER) for HSBC Mutual Fund"] + [None] * 13)
        wrong = list(amc_expenses._HSBC_HEADER)
        wrong[8] = "Calculated TER"
        ws.append(wrong)
        ws.append(self.row())
        out = io.BytesIO()
        wb.save(out)
        wb.close()
        with self.assertRaisesRegex(ValueError, "columns changed"):
            amc_expenses.parse_hsbc_workbook(out.getvalue(), date(2026, 9, 24))


class IciciExpenseTests(unittest.TestCase):
    def row(self, day="23/09/2026", *, scheme="ICICI Prudential Small Cap Fund",
            regular_ter="2.1%", direct_ter="1.18%"):
        return {
            "A": scheme, "B": day,
            "C": "1.5%", "D": "0.08%", "E": "0.01%", "F": "0.51%", "G": regular_ter,
            "H": "0.7%", "I": "0.08%", "J": "0.01%", "K": "0.39%", "L": direct_ter,
            "M": "NA", "N": "NA", "O": "NA", "P": "NA", "Q": "NA",
            "R": "NA", "S": "NA", "T": "NA", "U": "NA", "V": "NA",
        }

    def workbook(self, data_rows, *, header_override=None):
        rows = [
            {"A": "Total Expense Ratio (TER) for Mutual Fund Schemes"},
            {"C": "Regular Plan", "H": "Direct Plan", "M": "Unclaimed Dividend", "R": "Unclaimed Redemption"},
            dict(amc_expenses._ICICI_HEADER),
        ] + data_rows
        if header_override:
            rows[2].update(header_override)

        def xml_row(number, values):
            cells = []
            for column, value in values.items():
                text = escape(str(value))
                cells.append(
                    f'<c r="{column}{number}" t="inlineStr"><is><t>{text}</t></is></c>'
                )
            return f'<row r="{number}">' + "".join(cells) + "</row>"

        xml = (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<dimension ref="A1"/>'
            '<sheetData>'
            + "".join(xml_row(i + 1, row) for i, row in enumerate(rows))
            + "</sheetData></worksheet>"
        )
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/worksheets/sheet1.xml", xml)
        return out.getvalue()

    def categories(self):
        return [{
            "id": "parent",
            "internalName": "total-expense-ratio",
            "title": {"text": "Total Expense Ratio", "code": "TOTAL_EXPENSE_RATIO"},
            "isEnabled": True,
            "subCategory": [{
                "id": "child",
                "internalName": "Total Expense Ratio",
                "title": {"text": "Total Expense Ratio", "code": "TOTAL_EXPENSE_RATIO"},
                "isEnabled": True,
                "filter": [
                    {"key": {"code": "SHOW"}, "value": [{"code": "TER Details"}, {"code": "PDF"}]},
                    {"key": {"code": "FINANCIAL_YEAR"}, "value": [{"code": "2026-2027"}]},
                ],
            }],
        }]

    def file(self, month="September", short="Sep", *, parent="parent", child="child"):
        return {
            "title": {"text": f"TotalExpenseRatio{month}2026", "code": f"TotalExpenseRatio{month}2026"},
            "url": f"/financials-disclosures-files/Files/Total Expense Ratio/2026-2027/TotalExpenseRatio{short}2026.xlsx",
            "category": "TOTAL_EXPENSE_RATIO",
            "categoryName": "Total Expense Ratio",
            "isEnabled": True,
            "FINANCIAL_YEAR": ["2026-2027"],
            "SHOW": ["TER Details"],
            "level1Id": parent,
            "level2Id": child,
        }

    def test_icici_category_and_latest_file_identity_are_strict(self):
        parent, child, fy = amc_expenses._icici_category_ids(
            self.categories(), date(2026, 9, 24)
        )
        self.assertEqual((parent, child, fy), ("parent", "child", "2026-2027"))
        files = [self.file("September", "Sep"), self.file("August", "Aug")]
        source = amc_expenses._icici_select_file(
            files, parent, child, fy, date(2026, 9, 24)
        )
        self.assertEqual(
            source,
            "https://app.beta.icicipruamc.com/blob/financials-disclosures-files/Files/"
            "Total%20Expense%20Ratio/2026-2027/TotalExpenseRatioSep2026.xlsx",
        )

    def test_icici_future_file_and_wrong_metadata_cannot_be_promoted(self):
        future = self.file("October", "Oct")
        current = self.file()
        source = amc_expenses._icici_select_file(
            [future, current], "parent", "child", "2026-2027", date(2026, 9, 24)
        )
        self.assertTrue(source.endswith("TotalExpenseRatioSep2026.xlsx"))
        bad = self.file(parent="other")
        with self.assertRaisesRegex(ValueError, "no current TER Details"):
            amc_expenses._icici_select_file(
                [bad], "parent", "child", "2026-2027", date(2026, 9, 24)
            )

    def test_icici_latest_exact_row_retains_published_ber_and_total_ter(self):
        book = self.workbook([self.row("22/09/2026"), self.row("23/09/2026")])
        day, plans = amc_expenses.parse_icici_workbook(book, date(2026, 9, 24))
        self.assertEqual(day, "2026-09-23")
        self.assertEqual(plans["Regular"]["base_expense_ratio"], 1.5)
        self.assertEqual(plans["Regular"]["ter"], 2.1)
        self.assertEqual(plans["Direct"]["base_expense_ratio"], 0.7)
        self.assertEqual(plans["Direct"]["ter"], 1.18)

    def test_icici_future_row_wrong_scheme_and_duplicate_are_rejected(self):
        book = self.workbook([self.row("23/09/2026"), self.row("25/09/2026")])
        day, plans = amc_expenses.parse_icici_workbook(book, date(2026, 9, 24))
        self.assertEqual(day, "2026-09-23")
        self.assertEqual(plans["Direct"]["ter"], 1.18)

        wrong = self.workbook([self.row(scheme="ICICI Prudential MidCap Fund")])
        with self.assertRaisesRegex(ValueError, "no dated Small Cap rows"):
            amc_expenses.parse_icici_workbook(wrong, date(2026, 9, 24))

        row = self.row()
        duplicate = self.workbook([row, dict(row)])
        with self.assertRaisesRegex(ValueError, "duplicate Small Cap rows"):
            amc_expenses.parse_icici_workbook(duplicate, date(2026, 9, 24))

    def test_icici_header_and_component_reconciliation_are_strict(self):
        changed = self.workbook([self.row()], header_override={"G": "Calculated TER (%)"})
        with self.assertRaisesRegex(ValueError, "column G changed"):
            amc_expenses.parse_icici_workbook(changed, date(2026, 9, 24))

        bad_total = self.workbook([self.row(direct_ter="1.25%")])
        with self.assertRaisesRegex(ValueError, "components do not reconcile"):
            amc_expenses.parse_icici_workbook(bad_total, date(2026, 9, 24))

    @patch("tracker.amc_expenses.db.metric")
    @patch("tracker.amc_expenses.db.one", return_value={"code": "ICICI"})
    @patch("tracker.amc_expenses._icici_disclosure")
    def test_icici_collector_stores_exact_source_hash_and_plan_metrics(
        self, mock_disclosure, _mock_one, mock_metric
    ):
        source = (
            "https://app.beta.icicipruamc.com/blob/financials-disclosures-files/Files/"
            "Total%20Expense%20Ratio/2026-2027/TotalExpenseRatioSep2026.xlsx"
        )
        mock_disclosure.return_value = (source, self.workbook([self.row()]), "hash456")
        result = amc_expenses.icici(today=date(2026, 9, 24))
        self.assertIn("2026-09-23", result)
        self.assertEqual(mock_metric.call_count, 10)
        calls = {
            (call.args[1], call.args[2]): call.args[4]
            for call in mock_metric.call_args_list
        }
        self.assertEqual(calls[("Regular", "ter")], 2.1)
        self.assertEqual(calls[("Direct", "ter")], 1.18)
        self.assertEqual(calls[("Direct", "base_expense_ratio")], 0.7)
        self.assertTrue(all(call.args[3] == "2026-09-23" for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[6] == source for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[7] == "hash456" for call in mock_metric.call_args_list))


class InvescoExpenseTests(unittest.TestCase):
    def row(self, day="23/09/2026", *, scheme="Invesco India Small Cap Fund",
            nsdl="INVM/O/E/SCF/18/07/0030", regular_ter="1.83%", direct_ter="0.63%"):
        return {
            "Sr.No.": 23,
            "Scheme Name": scheme,
            "NSDL Scheme Code": nsdl,
            "TER Date(DD/MM/YYYY)": day,
            "Regular Plan - Base Expense Ratio (BER) (%)": "1.44%",
            "Regular Plan - Brokerage cost (%)": "0.05%",
            "Regular Plan - Transaction Cost incurred for the purpose of execution of trade (%)": "0.00%",
            "Regular Plan - Statutory Levies (including GST) (%)": "0.34%",
            "Regular Plan - Total TER (%)": regular_ter,
            "Direct Plan - Base Expense Ratio (BER) (%)": "0.41%",
            "Direct Plan - Brokerage cost (%)": "0.05%",
            "Direct Plan - Transaction Cost incurred for the purpose of execution of trade (%)": "0.00%",
            "Direct Plan - Statutory Levies (including GST) (%)": "0.17%",
            "Direct Plan - Total TER (%)": direct_ter,
        }

    def test_invesco_latest_exact_row_retains_published_ber_and_total_ter(self):
        rows = [self.row("22/09/2026"), self.row("23/09/2026")]
        day, plans = amc_expenses.parse_invesco_records(rows, date(2026, 9, 24))
        self.assertEqual(day, "2026-09-23")
        self.assertEqual(plans["Regular"]["base_expense_ratio"], 1.44)
        self.assertEqual(plans["Regular"]["ter"], 1.83)
        self.assertEqual(plans["Direct"]["base_expense_ratio"], 0.41)
        self.assertEqual(plans["Direct"]["ter"], 0.63)

    def test_invesco_future_wrong_identity_and_duplicate_are_rejected(self):
        rows = [self.row("23/09/2026"), self.row("25/09/2026")]
        day, plans = amc_expenses.parse_invesco_records(rows, date(2026, 9, 24))
        self.assertEqual(day, "2026-09-23")
        self.assertEqual(plans["Direct"]["ter"], 0.63)

        with self.assertRaisesRegex(ValueError, "no dated Small Cap rows"):
            amc_expenses.parse_invesco_records(
                [self.row(scheme="Invesco India Mid Cap Fund")], date(2026, 9, 24)
            )
        with self.assertRaisesRegex(ValueError, "no dated Small Cap rows"):
            amc_expenses.parse_invesco_records(
                [self.row(nsdl="INVM/O/E/OTHER")], date(2026, 9, 24)
            )

        row = self.row()
        with self.assertRaisesRegex(ValueError, "duplicate Small Cap rows"):
            amc_expenses.parse_invesco_records([row, dict(row)], date(2026, 9, 24))

    def test_invesco_component_reconciliation_is_strict(self):
        with self.assertRaisesRegex(ValueError, "components do not reconcile"):
            amc_expenses.parse_invesco_records(
                [self.row(direct_ter="0.70%")], date(2026, 9, 24)
            )

    def test_invesco_financial_year_and_month_fallback_roll_correctly(self):
        self.assertEqual(amc_expenses._invesco_financial_year_start(date(2026, 4, 1)), 2026)
        self.assertEqual(amc_expenses._invesco_financial_year_start(date(2026, 3, 31)), 2025)
        periods = amc_expenses._invesco_periods(date(2026, 4, 1))
        self.assertEqual(periods, (date(2026, 4, 1), date(2026, 3, 1)))

    @patch("tracker.amc_expenses.db.metric")
    @patch("tracker.amc_expenses.db.one", return_value={"code": "INVESCO"})
    @patch("tracker.amc_expenses._invesco_disclosure")
    def test_invesco_collector_stores_exact_source_hash_and_plan_metrics(
        self, mock_disclosure, _mock_one, mock_metric
    ):
        source = (
            "https://www.invescomutualfund.com/api/"
            "TotalExpenseRatioOfMutualFundSchemePolicy/GetTERExpenseData"
            "?title=Invesco+India+Small+Cap+Fund&fincialYear=2026&month=9"
        )
        _, plans = amc_expenses.parse_invesco_records(
            [self.row()], date(2026, 9, 24)
        )
        mock_disclosure.return_value = (source, "2026-09-23", plans, "hash789")
        result = amc_expenses.invesco(today=date(2026, 9, 24))
        self.assertIn("2026-09-23", result)
        self.assertEqual(mock_metric.call_count, 10)
        calls = {
            (call.args[1], call.args[2]): call.args[4]
            for call in mock_metric.call_args_list
        }
        self.assertEqual(calls[("Regular", "ter")], 1.83)
        self.assertEqual(calls[("Direct", "ter")], 0.63)
        self.assertEqual(calls[("Direct", "base_expense_ratio")], 0.41)
        self.assertTrue(all(call.args[3] == "2026-09-23" for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[6] == source for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[7] == "hash789" for call in mock_metric.call_args_list))


class JmExpenseTests(unittest.TestCase):
    def row(self, day="2026-09-24T00:00:00", *, scheme="JM Small Cap Fund",
            code="SC", nsdl="JMFI/O/E/SCF/23/11/0016",
            regular_ter=2.54, direct_ter=0.97):
        return {
            "NsdlSchemeCode": nsdl,
            "Scheme": scheme,
            "Schemecode": code,
            "TERDate": day,
            "RegularBER": 1.94,
            "RegularBrokCost": 0.09,
            "RegularTransCost": 0.01,
            "RegularStatLevGST": 0.50,
            "RegularTotalTER": regular_ter,
            "DirectBER": 0.59,
            "DirectBrokCost": 0.09,
            "DirectTransCost": 0.01,
            "DirectStatLevGST": 0.28,
            "DirectTotalTER": direct_ter,
        }

    def test_jm_latest_exact_row_retains_published_ber_and_total_ter(self):
        rows = [self.row("2026-09-23T00:00:00"), self.row()]
        day, plans = amc_expenses.parse_jm_ter_records(rows, date(2026, 9, 24))
        self.assertEqual(day, "2026-09-24")
        self.assertEqual(plans["Regular"]["base_expense_ratio"], 1.94)
        self.assertEqual(plans["Regular"]["ter"], 2.54)
        self.assertEqual(plans["Direct"]["base_expense_ratio"], 0.59)
        self.assertEqual(plans["Direct"]["ter"], 0.97)

    def test_jm_future_wrong_identity_and_duplicate_are_rejected(self):
        rows = [self.row(), self.row("2026-09-25T00:00:00")]
        day, plans = amc_expenses.parse_jm_ter_records(rows, date(2026, 9, 24))
        self.assertEqual(day, "2026-09-24")
        self.assertEqual(plans["Direct"]["ter"], 0.97)

        for kwargs in (
            {"scheme": "JM Midcap Fund"},
            {"code": "OTHER"},
            {"nsdl": "JMFI/O/E/OTHER"},
        ):
            with self.assertRaisesRegex(ValueError, "no dated Small Cap rows"):
                amc_expenses.parse_jm_ter_records(
                    [self.row(**kwargs)], date(2026, 9, 24)
                )

        row = self.row()
        with self.assertRaisesRegex(ValueError, "duplicate Small Cap rows"):
            amc_expenses.parse_jm_ter_records(
                [row, dict(row)], date(2026, 9, 24)
            )

    def test_jm_component_reconciliation_is_strict(self):
        with self.assertRaisesRegex(ValueError, "components do not reconcile"):
            amc_expenses.parse_jm_ter_records(
                [self.row(direct_ter=1.05)], date(2026, 9, 24)
            )

    @patch("tracker.amc_expenses.db.metric")
    @patch("tracker.amc_expenses.db.one", return_value={"code": "JM"})
    @patch("tracker.amc_expenses._jm_ter_disclosure")
    def test_jm_collector_stores_exact_source_hash_and_plan_metrics(
        self, mock_disclosure, _mock_one, mock_metric
    ):
        _, plans = amc_expenses.parse_jm_ter_records(
            [self.row()], date(2026, 9, 24)
        )
        mock_disclosure.return_value = ("2026-09-24", plans, "jmhash")
        result = amc_expenses.jm(today=date(2026, 9, 24))
        self.assertIn("2026-09-24", result)
        self.assertEqual(mock_metric.call_count, 10)
        calls = {
            (call.args[1], call.args[2]): call.args[4]
            for call in mock_metric.call_args_list
        }
        self.assertEqual(calls[("Regular", "ter")], 2.54)
        self.assertEqual(calls[("Direct", "ter")], 0.97)
        self.assertEqual(calls[("Direct", "base_expense_ratio")], 0.59)
        self.assertTrue(all(call.args[3] == "2026-09-24" for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[6] == amc_expenses.JM_TER_API for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[7] == "jmhash" for call in mock_metric.call_args_list))


class MahindraExpenseTests(unittest.TestCase):
    def workbook(self, rows, *, header_override=None):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        header1 = list(amc_expenses._MAHINDRA_HEADER_1)
        header2 = list(amc_expenses._MAHINDRA_HEADER_2)
        if header_override:
            header2[header_override[0]] = header_override[1]
        ws.append(header1)
        ws.append(header2)
        for row in rows:
            ws.append(row)
        out = io.BytesIO()
        wb.save(out)
        wb.close()
        return out.getvalue()

    def row(self, day="24-Sep-2026", *, scheme="Mahindra Manulife Small Cap Fund",
            nsdl="MAHM/O/E/SCF/22/07/0020", regular_ter=2.20, direct_ter=0.92):
        return [
            nsdl, scheme, day,
            1.59, 0.12, 0.01, 0.48, regular_ter,
            0.47, 0.12, 0.01, 0.32, direct_ter, None,
        ]

    def tree(self, *, title="TOTAL EXPENSE RATIO - 2026-27",
             url="https://www.mahindramanulife.com/uploads/download/current.xlsx"):
        return [{
            "categoryName": "MANDATORY DISCLOSURES",
            "subcategories": [{
                "categoryName": "Total Expense Ratio of Mutual Fund Schemes",
                "subcategories": [{
                    "categoryName": "Total Expense Ratio",
                    "files": [{"title": title, "fileUrl": url}],
                    "subcategories": [],
                }],
                "files": [],
            }],
            "files": [],
        }]

    def test_mahindra_current_financial_year_file_selection_is_exact(self):
        source = amc_expenses._mahindra_select_ter_file(
            self.tree(), date(2026, 9, 25)
        )
        self.assertEqual(
            source,
            "https://www.mahindramanulife.com/uploads/download/current.xlsx",
        )
        self.assertEqual(
            amc_expenses._mahindra_financial_year_title(date(2026, 3, 31)),
            "TOTAL EXPENSE RATIO - 2025-26",
        )
        self.assertEqual(
            amc_expenses._mahindra_financial_year_title(date(2026, 4, 1)),
            "TOTAL EXPENSE RATIO - 2026-27",
        )

    def test_mahindra_wrong_year_host_or_duplicate_file_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "not uniquely identified"):
            amc_expenses._mahindra_select_ter_file(
                self.tree(title="TOTAL EXPENSE RATIO - 2025-26"),
                date(2026, 9, 25),
            )
        with self.assertRaisesRegex(ValueError, "not uniquely identified"):
            amc_expenses._mahindra_select_ter_file(
                self.tree(url="https://example.com/uploads/download/current.xlsx"),
                date(2026, 9, 25),
            )
        tree = self.tree()
        tree[0]["subcategories"][0]["subcategories"][0]["files"].append(
            {
                "title": "TOTAL EXPENSE RATIO - 2026-27",
                "fileUrl": "https://www.mahindramanulife.com/uploads/download/other.xlsx",
            }
        )
        with self.assertRaisesRegex(ValueError, "not uniquely identified"):
            amc_expenses._mahindra_select_ter_file(tree, date(2026, 9, 25))

    def test_mahindra_latest_exact_row_retains_published_ber_and_total_ter(self):
        book = self.workbook([
            self.row("23-Sep-2026"),
            self.row("24-Sep-2026"),
        ])
        day, plans = amc_expenses.parse_mahindra_workbook(
            book, date(2026, 9, 25)
        )
        self.assertEqual(day, "2026-09-24")
        self.assertEqual(plans["Regular"]["base_expense_ratio"], 1.59)
        self.assertEqual(plans["Regular"]["ter"], 2.20)
        self.assertEqual(plans["Direct"]["base_expense_ratio"], 0.47)
        self.assertEqual(plans["Direct"]["ter"], 0.92)

    def test_mahindra_future_wrong_identity_and_duplicate_are_rejected(self):
        book = self.workbook([
            self.row("24-Sep-2026"),
            self.row("26-Sep-2026"),
        ])
        day, plans = amc_expenses.parse_mahindra_workbook(
            book, date(2026, 9, 25)
        )
        self.assertEqual(day, "2026-09-24")
        self.assertEqual(plans["Direct"]["ter"], 0.92)

        for kwargs in (
            {"scheme": "Mahindra Manulife Mid Cap Fund"},
            {"nsdl": "MAHM/O/E/OTHER"},
        ):
            bad = self.workbook([self.row(**kwargs)])
            with self.assertRaisesRegex(ValueError, "no dated Small Cap rows"):
                amc_expenses.parse_mahindra_workbook(
                    bad, date(2026, 9, 25)
                )

        row = self.row()
        duplicate = self.workbook([row, list(row)])
        with self.assertRaisesRegex(ValueError, "duplicate Small Cap rows"):
            amc_expenses.parse_mahindra_workbook(
                duplicate, date(2026, 9, 25)
            )

    def test_mahindra_header_and_component_reconciliation_are_strict(self):
        changed = self.workbook(
            [self.row()],
            header_override=(7, "Calculated TER (%)"),
        )
        with self.assertRaisesRegex(ValueError, "metric columns changed"):
            amc_expenses.parse_mahindra_workbook(
                changed, date(2026, 9, 25)
            )

        bad_total = self.workbook([self.row(direct_ter=1.00)])
        with self.assertRaisesRegex(ValueError, "components do not reconcile"):
            amc_expenses.parse_mahindra_workbook(
                bad_total, date(2026, 9, 25)
            )

    @patch("tracker.amc_expenses.db.metric")
    @patch("tracker.amc_expenses.db.one", return_value={"code": "MAHINDRA"})
    @patch("tracker.amc_expenses._mahindra_disclosure")
    def test_mahindra_collector_stores_exact_source_hash_and_plan_metrics(
        self, mock_disclosure, _mock_one, mock_metric
    ):
        source = "https://www.mahindramanulife.com/uploads/download/current.xlsx"
        _, plans = amc_expenses.parse_mahindra_workbook(
            self.workbook([self.row()]), date(2026, 9, 25)
        )
        mock_disclosure.return_value = (source, "2026-09-24", plans, "mahhash")
        result = amc_expenses.mahindra(today=date(2026, 9, 25))
        self.assertIn("2026-09-24", result)
        self.assertEqual(mock_metric.call_count, 10)
        calls = {
            (call.args[1], call.args[2]): call.args[4]
            for call in mock_metric.call_args_list
        }
        self.assertEqual(calls[("Regular", "ter")], 2.20)
        self.assertEqual(calls[("Direct", "ter")], 0.92)
        self.assertEqual(calls[("Direct", "base_expense_ratio")], 0.47)
        self.assertTrue(all(call.args[3] == "2026-09-24" for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[6] == source for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[7] == "mahhash" for call in mock_metric.call_args_list))


class MiraeExpenseTests(unittest.TestCase):
    def row(self, day=46288.0, *, scheme="Mirae Asset Small Cap Fund",
            nsdl="MIRA/O/E/SCF/24/10/0075", direct_ter=0.0064):
        return [
            nsdl, scheme, day,
            0.0157, 0.0007, 0.0001, 0.0047, 0.0212,
            0.0032, 0.0007, 0.0001, 0.0024, direct_ter,
        ]

    def test_latest_exact_mirae_row_retains_explicit_ber_and_total_ter(self):
        day, plans = amc_expenses.parse_mirae_rows(
            [self.row(46287.0), self.row(46288.0)],
            datemode=0,
            today=date(2026, 9, 23),
        )
        self.assertEqual(day, "2026-09-23")
        self.assertEqual(plans["Regular"]["base_expense_ratio"], 1.57)
        self.assertEqual(plans["Regular"]["ter"], 2.12)
        self.assertEqual(plans["Direct"]["base_expense_ratio"], 0.32)
        self.assertEqual(plans["Direct"]["ter"], 0.64)

    def test_mirae_identity_duplicate_and_reconciliation_are_strict(self):
        for kwargs in (
            {"scheme": "Mirae Asset Midcap Fund"},
            {"nsdl": "MIRA/O/E/OTHER"},
        ):
            with self.assertRaisesRegex(ValueError, "no dated Small Cap rows"):
                amc_expenses.parse_mirae_rows(
                    [self.row(**kwargs)], datemode=0, today=date(2026, 9, 23)
                )

        row = self.row()
        with self.assertRaisesRegex(ValueError, "duplicate Small Cap rows"):
            amc_expenses.parse_mirae_rows(
                [row, list(row)], datemode=0, today=date(2026, 9, 23)
            )

        with self.assertRaisesRegex(ValueError, "components do not reconcile"):
            amc_expenses.parse_mirae_rows(
                [self.row(direct_ter=0.0070)], datemode=0, today=date(2026, 9, 23)
            )

    def test_mirae_download_metadata_requires_matching_title_file_and_publish_date(self):
        payload = {
            "ReturnCode": "0",
            "Data": [
                {
                    "Title": "Total Expense Ratio -23 Sep 2026",
                    "URL": "/DailyUploads/TotalExpenseRatio/IN_MF_EXPENSE_RATIO_SEBI_V3_23092026.xls",
                    "PublishDate": "/Date(1790121600000)/",
                },
                {
                    "Title": "Total Expense Ratio -24 Sep 2026",
                    "URL": "/DailyUploads/TotalExpenseRatio/IN_MF_EXPENSE_RATIO_SEBI_V3_24092026.xls",
                    "PublishDate": "/Date(1790208000000)/",
                },
            ],
        }
        day, source = amc_expenses._mirae_select_download(
            payload, today=date(2026, 9, 23)
        )
        self.assertEqual(day, "2026-09-23")
        self.assertEqual(
            source,
            "https://www.miraeassetmf.co.in/DailyUploads/TotalExpenseRatio/"
            "IN_MF_EXPENSE_RATIO_SEBI_V3_23092026.xls",
        )

        payload["Data"][0]["PublishDate"] = "/Date(1790035200000)/"
        with self.assertRaisesRegex(ValueError, "no current valid workbook"):
            amc_expenses._mirae_select_download(
                {"ReturnCode": "0", "Data": [payload["Data"][0]]},
                today=date(2026, 9, 23),
            )

    @patch("tracker.amc_expenses.db.metric")
    @patch("tracker.amc_expenses.db.one", return_value={"code": "MIRAE"})
    @patch("tracker.amc_expenses._mirae_disclosure")
    def test_mirae_collector_stores_source_hash_and_all_reported_components(
        self, mock_disclosure, _mock_one, mock_metric
    ):
        day, plans = amc_expenses.parse_mirae_rows(
            [self.row()], datemode=0, today=date(2026, 9, 23)
        )
        source = (
            "https://www.miraeassetmf.co.in/DailyUploads/TotalExpenseRatio/"
            "IN_MF_EXPENSE_RATIO_SEBI_V3_23092026.xls"
        )
        mock_disclosure.return_value = (source, day, plans, "miraehash")
        result = amc_expenses.mirae(today=date(2026, 9, 24))
        self.assertIn("2026-09-23", result)
        self.assertEqual(mock_metric.call_count, 10)
        calls = {
            (call.args[1], call.args[2]): call.args[4]
            for call in mock_metric.call_args_list
        }
        self.assertEqual(calls[("Regular", "ter")], 2.12)
        self.assertEqual(calls[("Direct", "ter")], 0.64)
        self.assertEqual(calls[("Direct", "base_expense_ratio")], 0.32)
        self.assertTrue(all(call.args[3] == "2026-09-23" for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[6] == source for call in mock_metric.call_args_list))
        self.assertTrue(all(call.args[7] == "miraehash" for call in mock_metric.call_args_list))


if __name__ == "__main__":
    unittest.main()
