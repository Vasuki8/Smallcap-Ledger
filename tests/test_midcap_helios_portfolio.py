"""Helios discovery and workbook boundaries, with real in-memory XLSX parsing."""
from datetime import date, datetime
import hashlib
import io
import unittest
from unittest.mock import patch

import openpyxl

from tracker.midcap_portfolio_structured import collect


PAGE = "https://www.heliosmf.in/portfolio-disclosure"
SOURCE = "https://www.heliosmf.in/wp-content/uploads/2026/09/helios-mid-cap-fund-monthly-portfolio-as-on-31st-august-2026.xlsx"
FAMILY = "Helios Mid Cap Fund"


def page(source=SOURCE, label="August 2026", family=FAMILY, category="Monthly Portfolio", extra=""):
    return f'''<div class="hlx-dl-cat-item" data-depth="1">
    <h3>{category}</h3><div class="hlx-dl-cat-body">
    <div class="hlx-dl-cat-item" data-depth="2"><h4>{family}</h4>
    <div class="hlx-dl-cat-body"><div class="hlx-dl-cat-item" data-depth="3">
    <h5>2026</h5><div class="hlx-dl-cat-body"><div class="hlx-dl-files-grid">
    <a class="hlx-dl-file" href="{source}">{label}</a>{extra}
    </div></div></div></div></div></div></div>'''.encode()


def workbook(family=FAMILY, day=datetime(2026, 8, 31)):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "HMCF"
    sheet.append([None, None, "Helios Mutual Fund"])
    sheet.append([None, None, "SCHEME NAME :", family + " (An open-ended equity scheme predominantly investing in mid cap stocks)"])
    sheet.append([None, None, "PORTFOLIO STATEMENT AS ON :", day])
    sheet.append([None, None, "Name of the Instrument / Issuer", "ISIN", "Rating / Industry^", "Quantity", "Market value (Rs. in Lakhs)", "% to AUM"])
    sheet.append([None, None, "EQUITY & EQUITY RELATED"])
    for name, isin in [
        ("Alpha Ltd.", "INE000A01010"), ("Beta Ltd.", "INE000B01010"),
        ("Gamma Ltd.", "INE000C01010"), ("Delta Ltd.", "INE000D01010"),
        ("Epsilon Ltd.", "INE000E01010"),
    ]:
        sheet.append([None, None, name, isin, "Industry", 100, 180, 18])
    sheet.append([None, None, "Net Current Assets", None, None, None, 100, 10])
    sheet.append([None, None, "Grand Total", None, None, None, 1000, 100])
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


class HeliosPortfolioTests(unittest.TestCase):
    def inspect(self, html=None, content=None, media="text/html; charset=utf-8"):
        from tracker.midcap_helios_portfolio import inspect
        html = html if html is not None else page()
        content = content if content is not None else workbook()
        def fetch(url, **kwargs):
            self.assertFalse(kwargs["archive"])
            if url == PAGE:
                return html, None, media
            self.assertEqual(url, SOURCE)
            return content, None, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        with patch("tracker.db.connect", side_effect=AssertionError("No production DB writes")):
            return inspect("2026-08-31", fetch)

    def test_dated_exact_scheme_workbook_preserves_complete_evidence_and_provenance(self):
        content = workbook()
        row = self.inspect(content=content)
        self.assertEqual(row["family"], FAMILY)
        self.assertEqual(row["as_of"], "2026-08-31")
        self.assertEqual(row["positions_observed"], 6)
        self.assertTrue(row["complete"])
        self.assertEqual(row["aum"], 10)
        self.assertEqual(row["source"], SOURCE)
        self.assertEqual(row["source_sha256"], hashlib.sha256(content).hexdigest())
        self.assertEqual(row["discovery_source_sha256"], hashlib.sha256(page()).hexdigest())
        self.assertEqual(row["sheet"], "HMCF")

    def test_wrong_scheme_and_half_yearly_sections_cannot_supply_monthly_evidence(self):
        for html in (page(family="Helios Large & Mid Cap Fund"), page(category="Half Yearly Portfolio")):
            with self.subTest(html=html), self.assertRaises(ValueError):
                self.inspect(html=html)

    def test_stale_label_or_full_date_in_filename_is_not_current(self):
        for html in (page(label="July 2026"), page(source=SOURCE.replace("31st-august", "30th-august"))):
            with self.subTest(html=html), self.assertRaises(ValueError):
                self.inspect(html=html)

    def test_multiple_distinct_current_links_fail_closed(self):
        extra = '<a class="hlx-dl-file" href="' + SOURCE.replace("/2026/09/", "/2026/10/") + '">August 2026</a>'
        with self.assertRaises(ValueError):
            self.inspect(html=page(extra=extra))

    def test_identical_repeated_link_is_one_source(self):
        extra = f'<a class="hlx-dl-file" href="{SOURCE}">August 2026</a>'
        self.assertEqual(self.inspect(html=page(extra=extra))["positions_observed"], 6)

    def test_unapproved_or_credentialled_download_url_fails_closed(self):
        for source in (SOURCE.replace("www.heliosmf.in", "evil.example"),
                       SOURCE.replace("https://", "http://"),
                       SOURCE.replace("https://", "https://user@"),
                       SOURCE.replace("https://", "https://@"),
                       SOURCE.replace("https://", "https://:@"),
                       SOURCE.replace("www.heliosmf.in", "www.heliosmf.in:444")):
            with self.subTest(source=source), self.assertRaises(ValueError):
                self.inspect(html=page(source=source))

    def test_wrong_workbook_family_or_reporting_date_is_rejected(self):
        for content in (workbook(family="Helios Large & Mid Cap Fund"),
                        workbook(day=datetime(2026, 7, 31))):
            with self.subTest(content=len(content)), self.assertRaises(ValueError):
                self.inspect(content=content)

    def test_unrecognized_numeric_row_keeps_portfolio_partial(self):
        book = openpyxl.load_workbook(io.BytesIO(workbook()))
        book.active.insert_rows(11)
        for col, value in enumerate([None, None, "Unknown Asset", None, None, None, 10, 1], 1):
            book.active.cell(11, col, value)
        out = io.BytesIO()
        book.save(out)
        book.close()
        row = self.inspect(content=out.getvalue())
        self.assertFalse(row["complete"])
        self.assertIn("Unknown Asset", row["unknown_rows"])

    def test_non_html_discovery_is_rejected(self):
        with self.assertRaises(ValueError):
            self.inspect(media="application/json")

    def test_existing_batch_collects_staged_helios_without_production_writes(self):
        content = workbook()
        def fetch(url, **kwargs):
            self.assertFalse(kwargs["archive"])
            if url == PAGE:
                return page(), None, "text/html"
            if url == SOURCE:
                return content, None, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            raise AssertionError("Unexpected source")
        with patch("tracker.db.rows", return_value=[{"family": FAMILY}]), \
             patch("tracker.db.connect", side_effect=AssertionError("No production DB writes")):
            report = collect(fetch_fn=fetch, today=date(2026, 9, 30))
        rows = [row for row in report["results"] if row["family"] == FAMILY]
        self.assertEqual(len(rows), 1)
        self.assertTrue(rows[0]["complete"])
        self.assertIn("observed_at", rows[0])
        self.assertEqual(report["production_writes"], 0)
        self.assertFalse(report["public_export_enabled"])


if __name__ == "__main__":
    unittest.main()
