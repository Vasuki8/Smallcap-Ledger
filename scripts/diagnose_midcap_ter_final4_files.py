"""Inspect four concrete official TER files/exports for exact staged Mid Cap identities."""
from __future__ import annotations

import argparse,io,json,re,hashlib
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))

import openpyxl,xlrd
from tracker import providers


TARGETS={
  "Bank of India Mid Cap Fund":"https://www.boimf.in/docs/default-source/investorcorner/total-expense-ratio/expense_ratio_01092026_to_30092026.xls?sfvrsn=1b375908_8",
  "Helios Mid Cap Fund":"https://www.heliosmf.in/wp-content/uploads/2026/09/TER_sep_2026_vsvefi.xls",
  "Kotak Mid Cap Fund":"https://vatseelabs-s3.kotakmf.com/TER/TER-2026-2027.xlsx",
  "The Wealth Company Mid Cap Fund":"https://www.wealthcompanyamc.in/api/ter-values/year/2026/export",
}


def compact(value):
    return re.sub(r"\s+"," ",str(value or "")).strip()


def match_row(row,family):
    return family.casefold().replace(" ","") in " ".join(compact(x) for x in row).casefold().replace(" ","")


def inspect_xlsx(body,family):
    book=openpyxl.load_workbook(io.BytesIO(body),data_only=True,read_only=True)
    try:
        out={"format":"xlsx","sheets":[]}
        for name in book.sheetnames:
            sh=book[name]
            rows=[]
            for i,row in enumerate(sh.iter_rows(values_only=True),1):
                vals=list(row)
                if i<=6 or match_row(vals,family):
                    rows.append({"row":i,"values":[compact(v) for v in vals[:24]]})
                if i>5000:break
            out["sheets"].append({"name":name,"max_row":sh.max_row,"max_column":sh.max_column,"samples":rows[:80]})
        return out
    finally:book.close()


def inspect_xls(body,family):
    book=xlrd.open_workbook(file_contents=body)
    try:
        out={"format":"xls","sheets":[]}
        for name in book.sheet_names():
            sh=book.sheet_by_name(name);rows=[]
            for i in range(min(sh.nrows,5000)):
                vals=sh.row_values(i)
                if i<6 or match_row(vals,family):
                    rows.append({"row":i+1,"values":[compact(v) for v in vals[:24]]})
            out["sheets"].append({"name":name,"max_row":sh.nrows,"max_column":sh.ncols,"samples":rows[:80]})
        return out
    finally:book.release_resources()


def inspect_json(body,family):
    payload=json.loads(body)
    text=json.dumps(payload,ensure_ascii=False)
    samples=[]
    def walk(node,path="$"):
        if len(samples)>=40:return
        if isinstance(node,dict):
            joined=" ".join(compact(v) for v in node.values() if not isinstance(v,(dict,list)))
            if family.casefold().replace(" ","") in joined.casefold().replace(" ",""):
                samples.append({"path":path,"row":node})
            for k,v in node.items():walk(v,path+"."+str(k))
        elif isinstance(node,list):
            for i,v in enumerate(node[:10000]):walk(v,f"{path}[{i}]")
    walk(payload)
    return {"format":"json","top_type":type(payload).__name__,"bytes":len(body),"matched_samples":samples}


def inspect(url,family):
    body,_,typ=providers.fetch(url,archive=False,max_bytes=25*1024*1024)
    result={"url":url,"bytes":len(body),"content_type":typ,"sha256":hashlib.sha256(body).hexdigest()}
    if body.startswith(b"PK"):
        result.update(inspect_xlsx(body,family))
    elif body.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        result.update(inspect_xls(body,family))
    else:
        try:result.update(inspect_json(body,family))
        except Exception:
            result["format"]="text"
            text=body.decode("utf-8","replace")
            result["snippets"]=[m.group(0) for m in re.finditer(r".{0,180}mid.?cap.{0,260}",text,re.I|re.S)][:20]
    return result


def main(argv=None):
    p=argparse.ArgumentParser()
    p.add_argument("--output",type=Path,required=True)
    args=p.parse_args(argv)
    out={"mode":"read_only","production_writes":0,"targets":[]}
    for family,url in TARGETS.items():
        try:r=inspect(url,family);r["family"]=family;r["status"]="ok"
        except Exception as exc:r={"family":family,"url":url,"status":"error","error":(str(exc) or type(exc).__name__)[:500]}
        out["targets"].append(r)
    out["summary"]={"ok":sum(x["status"]=="ok" for x in out["targets"]),"error":sum(x["status"]=="error" for x in out["targets"])}
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(out,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(out["summary"],indent=2))

if __name__=="__main__":main()
