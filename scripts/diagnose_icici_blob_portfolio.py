"""Read-only ICICI monthly portfolio transport probe using current blob delivery."""
from __future__ import annotations
import io,sys,zipfile
from pathlib import Path,PurePosixPath
from urllib.parse import urlparse
import httpx

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))

OLD=('https://www.icicipruamc.com/downloads/Files/Monthly%20Portfolio%20Disclosures/'
     '2026/Aug/Monthly-Portfolio-Disclosure-August-2026.zip')
BLOB=('https://www.icicipruamc.com/blob/downloads/Files/Monthly%20Portfolio%20Disclosures/'
      '2026/Aug/Monthly-Portfolio-Disclosure-August-2026.zip')

def safe(url):
    host=(urlparse(url).hostname or '').lower()
    return host=='www.icicipruamc.com' or host.endswith('.icicipruamc.com')

def probe(url,follow):
    if not safe(url):raise ValueError('Refusing non-ICICI host')
    with httpx.Client(
        timeout=httpx.Timeout(30,read=60),follow_redirects=follow,
        headers={'User-Agent':'Mozilla/5.0','Accept':'*/*',
                 'Referer':'https://www.icicipruamc.com/'}) as client:
        r=client.get(url)
        print('ICICI_BLOB_HTTP url='+url+' follow='+str(follow).lower()+
              f' status={r.status_code} bytes={len(r.content)} content_type={r.headers.get("content-type","")}'+
              ' location='+str(r.headers.get('location') or '')+' final='+str(r.url),flush=True)
        return r

def main():
    try:
        old=probe(OLD,False)
        location=old.headers.get('location')
        if location and location.startswith('https://') and safe(location):
            try:probe(location,True)
            except Exception as exc:
                print('ICICI_REDIRECT_ERROR '+(str(exc) or type(exc).__name__).splitlines()[0][:700],flush=True)
    except Exception as exc:
        print('ICICI_OLD_ERROR '+(str(exc) or type(exc).__name__).splitlines()[0][:700],flush=True)
    try:
        r=probe(BLOB,True)
        print('ICICI_BLOB_SIGNATURE '+r.content[:32].hex(),flush=True)
        if r.status_code==200 and r.content.startswith(b'PK'):
            print('ICICI_BLOB_ZIP_OK bytes='+str(len(r.content)),flush=True)
            with zipfile.ZipFile(io.BytesIO(r.content)) as z:
                entries=z.infolist()
                print(f'ICICI_ZIP_ENTRIES count={len(entries)} total_uncompressed={sum(x.file_size for x in entries)}',flush=True)
                spreadsheets=[]
                for entry in entries:
                    p=PurePosixPath(entry.filename)
                    safe=not (p.is_absolute() or '..' in p.parts or '\\' in entry.filename or entry.flag_bits&1)
                    print('ICICI_ZIP_MEMBER '+entry.filename+
                          f' bytes={entry.file_size} compressed={entry.compress_size} safe={safe}',flush=True)
                    if safe and p.suffix.lower() in ('.xls','.xlsx') and 0<entry.file_size<=40*1024*1024:
                        spreadsheets.append(entry)
                print(f'ICICI_ZIP_SPREADSHEETS count={len(spreadsheets)}',flush=True)
                import openpyxl,xlrd
                for entry in spreadsheets[:80]:
                    with z.open(entry) as fh:body=fh.read()
                    try:
                        if body.startswith(b'PK'):
                            book=openpyxl.load_workbook(io.BytesIO(body),read_only=True,data_only=True)
                            sheet_names=book.sheetnames
                            print('ICICI_WORKBOOK '+entry.filename+' sheets='+repr(sheet_names),flush=True)
                            for sh in book.worksheets:
                                if 'small' in sh.title.lower():
                                    rows=[]
                                    for row in sh.iter_rows(min_row=1,max_row=18,values_only=True):
                                        rows.append([str(x)[:160] if x is not None else '' for x in row[:12]])
                                    print('ICICI_SMALL_SHEET '+entry.filename+' :: '+sh.title+' :: '+repr(rows),flush=True)
                            book.close()
                        elif body.startswith(b'\xd0\xcf'):
                            book=xlrd.open_workbook(file_contents=body,on_demand=True)
                            print('ICICI_WORKBOOK '+entry.filename+' sheets='+repr(book.sheet_names()),flush=True)
                            for name in book.sheet_names():
                                if 'small' not in name.lower():continue
                                sh=book.sheet_by_name(name)
                                rows=[sh.row_values(i)[:12] for i in range(min(18,sh.nrows))]
                                print('ICICI_SMALL_SHEET '+entry.filename+' :: '+name+' :: '+repr(rows),flush=True)
                            book.release_resources()
                    except Exception as exc:
                        print('ICICI_WORKBOOK_ERROR '+entry.filename+' '+(str(exc) or type(exc).__name__).splitlines()[0][:500],flush=True)
    except Exception as exc:
        print('ICICI_BLOB_ERROR '+(str(exc) or type(exc).__name__).splitlines()[0][:700],flush=True)

if __name__=='__main__':main()
