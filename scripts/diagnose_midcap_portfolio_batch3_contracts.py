"""Discover exact first-party contracts for staged Mid Cap portfolio batch 3."""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
import zipfile
from pathlib import Path, PurePosixPath
from urllib.parse import urljoin

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))

import openpyxl
from bs4 import BeautifulSoup

from tracker import axis_portfolios, icici_portfolios, providers, sbi_portfolios


def _text(value):
    return re.sub(r"\s+"," ",str(value or "")).strip()


def _zip_entries(body, pattern=r"mid\s*cap|midcap"):
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        entries=archive.infolist()
        if len(entries)>250 or sum(x.file_size for x in entries)>150*1024*1024:
            raise ValueError("ZIP exceeds diagnostic safety bounds")
        out=[]
        for entry in entries:
            path=PurePosixPath(entry.filename)
            if path.is_absolute() or ".." in path.parts or "\\" in entry.filename or entry.flag_bits&1:
                raise ValueError("Unsafe ZIP entry")
            if re.search(pattern,entry.filename,re.I):
                out.append({"name":entry.filename,"bytes":entry.file_size})
        return out


def icici():
    candidates=list(icici_portfolios.discover(providers.fetch))
    current=candidates[0]
    _,source,title=current
    body,_,typ=providers.fetch(source,archive=False,max_bytes=100*1024*1024)
    return {
        "family":"ICICI Prudential Mid Cap Fund",
        "listing_candidates":[{"family":f,"url":u,"title":t} for f,u,t in candidates],
        "current_source":source,
        "current_title":title,
        "content_type":typ,
        "midcap_zip_entries":_zip_entries(body),
    }


def axis():
    token=axis_portfolios.cms_token(providers.fetch)
    headers={"Authorization":token}
    nested,_,_=providers.fetch(
        axis_portfolios.NESTED_ENDPOINT,
        body={"sdParentID":axis_portfolios.PARENT_ID},
        headers=headers,
        archive=False,
        max_bytes=4*1024*1024,
    )
    axis_portfolios.validate_monthly_branch(nested)
    docs,_,_=providers.fetch(
        axis_portfolios.DOCUMENTS_ENDPOINT,
        body={"sdType":axis_portfolios.MONTHLY_TYPE,"sdID":axis_portfolios.MONTHLY_ID},
        headers=headers,
        archive=False,
        max_bytes=8*1024*1024,
    )
    payload=axis_portfolios._json(docs,"monthly scheme documents endpoint")
    schemes=[
        {"schemeName":_text(x.get("schemeName")),"schemeCode":_text(x.get("schemeCode"))}
        for x in payload.get("schemeCategories") or []
        if isinstance(x,dict) and re.search(r"mid\s*cap|midcap",_text(x.get("schemeName")),re.I)
    ]
    documents=[]
    for item in payload.get("documentList") or []:
        if not isinstance(item,dict):continue
        title=_text(item.get("documentName"))
        target=_text(item.get("docuementURL"))
        if re.search(r"mid\s*cap|midcap",title+" "+target,re.I):
            documents.append({
                "documentName":title,
                "documentPostedDate":_text(item.get("documentPostedDate")),
                "url":target,
                "redirectionUrl":item.get("redirectionUrl"),
            })
    return {
        "family":"Axis Midcap Fund",
        "scheme_categories":schemes,
        "documents":documents[:50],
    }


def sbi():
    payload={"FundId":0,"PSYear":"2026","PSMonth":"August","PSFrequency":"Monthly"}
    raw,_,_=providers.fetch(
        sbi_portfolios.ENDPOINT,body=payload,archive=False,max_bytes=8*1024*1024)
    soup=BeautifulSoup(raw,"html.parser")
    rows=[]
    for anchor in soup.select("a[href]"):
        title=_text(anchor.get_text(" ",strip=True))
        href=urljoin(sbi_portfolios.ENDPOINT,anchor.get("href"))
        if re.search(r"mid\s*cap|midcap",title+" "+href,re.I):
            rows.append({"title":title,"url":href})
    return {"family":"SBI MIDCAP FUND","documents":rows[:50]}


def tata():
    api="https://prod-dist-api.tatamfdev.com/cms-data/api/CMSDATA_portfolio?type=monthly"
    raw,_,_=providers.fetch(api,archive=False,max_bytes=8*1024*1024)
    data=json.loads(raw)
    rows=[]
    for item in data if isinstance(data,list) else []:
        if not isinstance(item,dict):continue
        title=_text(item.get("field_document_title"))
        target=_text(item.get("field_media_document"))
        if re.search(r"31(?:st)?\s+August,?\s+2026|31\s+Aug(?:ust)?\s+2026",title,re.I):
            rows.append({"title":title,"url":target})
    if not rows:
        raise ValueError("Tata CMS exposed no August 31, 2026 monthly portfolio workbook")
    rows=rows[:4]
    inspected=[]
    for row in rows:
        body,_,typ=providers.fetch(row["url"],archive=False,max_bytes=30*1024*1024)
        if not body.startswith(b"PK"):
            inspected.append({**row,"error":"not XLSX","content_type":typ});continue
        book=openpyxl.load_workbook(io.BytesIO(body),read_only=True,data_only=True)
        matches=[]
        try:
            for sh in book.worksheets:
                prefix=[]
                for values in sh.iter_rows(min_row=1,max_row=min(30,sh.max_row),values_only=True):
                    prefix.extend(_text(v) for v in values if v is not None)
                joined=" ".join(prefix)
                if re.search(r"tata\s+mid\s*cap\s+fund",joined,re.I):
                    matches.append({"sheet":sh.title,"prefix":joined[:1800]})
        finally:
            book.close()
        inspected.append({**row,"content_type":typ,"midcap_sheets":matches})
    return {"family":"Tata Mid Cap Fund","api":api,"workbooks":inspected}


def uti():
    url="https://www.utimf.com/api/get-consolidate-portfolio-disclosure?year=2026&month=August"
    raw,_,_=providers.fetch(url,archive=False,max_bytes=8*1024*1024)
    data=json.loads(raw)
    rows=data.get("rows") or []
    out=[]
    for row in rows[:10]:
        if not isinstance(row,dict):continue
        target=_text(row.get("url"))
        name=_text(row.get("name"))
        typ=_text(row.get("type"))
        item={"name":name,"url":target,"type":typ}
        if typ.lower()=="zip" and target:
            body,_,media=providers.fetch(target,archive=False,max_bytes=100*1024*1024)
            item["content_type"]=media
            item["midcap_zip_entries"]=_zip_entries(body)
        out.append(item)
    return {"family":"UTI - Mid Cap Fund","api":url,"rows":out}


def run():
    result={"mode":"read_only","production_writes":0,"public_export_enabled":False,"targets":[]}
    for name,fn in (("ICICI",icici),("Axis",axis),("SBI",sbi),("Tata",tata),("UTI",uti)):
        try:
            item=fn();item["amc"]=name;item["status"]="ok"
        except Exception as exc:
            item={"amc":name,"status":"error","error":(str(exc) or type(exc).__name__)[:600]}
        result["targets"].append(item)
    result["summary"]={
        "targets":5,
        "ok":sum(x["status"]=="ok" for x in result["targets"]),
        "errors":sum(x["status"]=="error" for x in result["targets"]),
    }
    return result


def main(argv=None):
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--output",type=Path,required=True)
    args=p.parse_args(argv)
    result=run()
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(result["summary"],indent=2))


if __name__=="__main__":
    main()
