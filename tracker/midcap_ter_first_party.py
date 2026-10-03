"""Read-only first-party TER discovery for staged Mid Cap funds with reusable AMC transports."""
from __future__ import annotations

from collections import defaultdict
from datetime import date,timedelta,datetime
import calendar
import hashlib
import io
import json
import math
from urllib.parse import urlencode

import openpyxl
import xlrd

from . import amc_expenses,amfi_metrics,db,providers
from . import jm_portfolios


TARGETS={
    "BANK OF INDIA MID CAP FUND":"Bank of India Mutual Fund",
    "Canara Robeco Mid Cap Fund":"Canara Robeco Mutual Fund",
    "HSBC Midcap Fund":"HSBC Mutual Fund",
    "ICICI Prudential Mid Cap Fund":"ICICI Prudential Mutual Fund",
    "Invesco India Mid Cap Fund":"Invesco Mutual Fund",
    "JM Mid Cap Fund":"JM Financial Mutual Fund",
    "Mahindra Manulife Mid Cap Fund":"Mahindra Manulife Mutual Fund",
    "Mirae Asset Midcap Fund":"Mirae Asset Mutual Fund",
}


_BOI_SOURCE="https://www.boimf.in/docs/default-source/investorcorner/total-expense-ratio/expense_ratio_01092026_to_30092026.xls?sfvrsn=1b375908_8"
_BOI_PERIOD_START=date(2026,9,1)
_BOI_PERIOD_END=date(2026,9,30)
_BOI_NSDL="BOIA/O/E/MIF/25/06/0023"
_BOI_HEADER=(
    "NSDL Scheme Code","Scheme Name","TER Date\n(DD/MM/\nYYYY)",
    "Regular Plan - Base Expense Ratio (BER) (%)","Regular Plan - Brokerage cost (%)",
    "Regular Plan - Transaction Cost incurred for the purpose of execution of trade (%)",
    "Regular Plan - Statutory Levies (including GST) (%)","Regular Plan - Total TER (%)",
    "Direct Plan - Base Expense Ratio (BER) (%)","Direct Plan - Brokerage cost (%)",
    "Direct Plan - Transaction Cost incurred for the purpose of execution of trade (%)",
    "Direct Plan - Statutory Levies (including GST) (%)","Direct Plan - Total TER (%)",
)


def _same_family(published,staged):
    return amfi_metrics.normalized(published)==amfi_metrics.normalized(staged)


def _sha(content):
    return hashlib.sha256(content).hexdigest()


class _IciciWorkbookError(ValueError):
    """Keep rejected source identity separate from accepted TER evidence."""

    def __init__(self,message,source=None,content=None,*,api_error=None,rejected_candidates=None):
        super().__init__(message)
        self.source_evidence={
            "discovery_channel":"financial_disclosure_api" if api_error is None else "reviewed_monthly_workbook_fallback",
            "discovery_api_error":api_error,
        }
        if content is not None:
            self.source_evidence.update({"source":source,"sha256":_sha(content),"bytes":len(content)})
        if rejected_candidates:
            self.source_evidence["rejected_candidates"]=rejected_candidates


class _BoiWorkbookError(ValueError):
    """Retain actual acquired BOI bytes without inventing transport evidence."""

    def __init__(self,message,content=None):
        super().__init__(message)
        self.source_evidence={"source":_BOI_SOURCE,"discovery_channel":"reviewed_monthly_workbook"}
        if content is not None:
            self.source_evidence.update({"sha256":_sha(content),"bytes":len(content)})


def _reconcile(plans,label):
    for plan in ("Regular","Direct"):
        values=plans.get(plan)
        if not values:
            raise ValueError(f"{label} does not expose both Regular and Direct")
        if values["ter"]+1e-9<values["base_expense_ratio"]:
            raise ValueError(f"{label} {plan} Total TER is below BER")
        components=sum(values[k] for k in (
            "base_expense_ratio","brokerage","transaction_cost","statutory_levies"
        ))
        if abs(components-values["ter"])>0.02:
            raise ValueError(f"{label} {plan} TER components do not reconcile")
    return plans


def _latest_exact(matches,label):
    if not matches:
        raise ValueError(f"{label} contains no exact Mid Cap rows")
    day=max(matches)
    if len(matches[day])!=1:
        raise ValueError(f"{label} has duplicate exact rows for {day}")
    return day,matches[day][0]


def _boi_number(cell,label):
    value=cell.value
    if value is None or isinstance(value,bool) or str(value).strip().upper() in ("","-","NA","N/A"):
        raise ValueError(f"BOI TER workbook is missing or has invalid {label}")
    if "%" in cell.number_format:
        raise ValueError(f"BOI TER workbook {label} percentage format changed units")
    try:
        parsed=float(value)
    except (ValueError,TypeError) as exc:
        raise ValueError(f"BOI TER workbook has invalid numeric {label}") from exc
    if not math.isfinite(parsed) or not 0<=parsed<=5:
        raise ValueError(f"BOI TER workbook {label} is outside accepted numeric range")
    return parsed


def _boi(family,today):
    content=None
    try:
        month=today.replace(day=1)
        previous=(month-timedelta(days=1)).replace(day=1)
        if _BOI_PERIOD_START not in (month,previous):
            raise ValueError("BOI reviewed TER source period is outside current/previous month")
        content,_,_=providers.fetch(_BOI_SOURCE,archive=False,max_bytes=5*1024*1024)
        if not content.startswith(b"PK"):
            raise ValueError("BOI TER source is not XLSX")
        book=openpyxl.load_workbook(io.BytesIO(content),data_only=True,read_only=True)
        try:
            if book.sheetnames!=["TER_UPLOAD_FORMAT"]:
                raise ValueError("BOI TER workbook sheet layout changed")
            sheet=book["TER_UPLOAD_FORMAT"]
            # The reviewed workbook declares A1 while containing 713 rows.
            sheet.reset_dimensions()
            rows=sheet.iter_rows(max_col=13)
            header=next(rows,None)
            if tuple(str(c.value or "").strip() for c in (header or ()))!=_BOI_HEADER:
                raise ValueError("BOI TER workbook columns changed")
            matches=defaultdict(list)
            for row in rows:
                published=str(row[1].value or "").strip()
                nsdl=str(row[0].value or "").strip()
                same_family=_same_family(published,family)
                if not same_family and nsdl!=_BOI_NSDL:
                    continue
                if not same_family or nsdl!=_BOI_NSDL:
                    raise ValueError("BOI Mid Cap family/NSDL identity mismatch")
                raw=row[2].value
                if isinstance(raw,datetime):
                    day=raw.date()
                elif isinstance(raw,date):
                    day=raw
                else:
                    raise ValueError("BOI exact Mid Cap row is missing a valid styled date")
                if not _BOI_PERIOD_START<=day<=_BOI_PERIOD_END or day>today:
                    raise ValueError("BOI exact Mid Cap date is outside source period or in the future")
                matches[day.isoformat()].append(row)
            day,row=_latest_exact(matches,"BOI TER workbook")
            fields=("base_expense_ratio","brokerage","transaction_cost","statutory_levies","ter")
            plans={
                plan:{key:_boi_number(row[start+i],f"{plan} {key}") for i,key in enumerate(fields)}
                for plan,start in (("Regular",3),("Direct",8))
            }
            _reconcile(plans,"BOI Mid Cap")
            return {
                "family":family,"amc":TARGETS[family],"status":"recovered","as_of":day,
                "plans":plans,"source":_BOI_SOURCE,"sha256":_sha(content),"bytes":len(content),
                "identity":{"scheme_name":str(row[1].value).strip(),"nsdl_scheme_code":str(row[0].value).strip()},
                "evidence_type":"official_amc_reviewed_workbook","discovery_channel":"reviewed_monthly_workbook",
            }
        finally:
            book.close()
    except Exception as exc:
        raise _BoiWorkbookError(str(exc) or type(exc).__name__,content) from exc


def _canara(family,today):
    start=today-timedelta(days=14)
    url=amc_expenses.CANARA_API+"?"+urlencode({
        "from_date":start.isoformat(),"to_date":today.isoformat()
    })
    body,_,_=providers.fetch(url,archive=False,max_bytes=2*1024*1024)
    records=json.loads(body)
    if not isinstance(records,list):
        raise ValueError("Canara TER API response changed format")
    grouped=defaultdict(list)
    identities=set()
    for row in records:
        if not isinstance(row,dict) or not _same_family(row.get("scheme_name",""),family):
            continue
        day=str(row.get("date") or "").strip()
        try:parsed=date.fromisoformat(day)
        except ValueError:continue
        if parsed>today:continue
        plan=str(row.get("plan_type") or "").strip()
        if plan not in amc_expenses.CANARA_PLANS:continue
        identities.add((str(row.get("scheme_name") or "").strip(),str(row.get("sch_code") or "").strip()))
        grouped[day].append(row)
    if len(identities)!=1:
        raise ValueError("Canara exact Mid Cap identity is not unique")
    for day in sorted(grouped,reverse=True):
        by_plan=defaultdict(list)
        for row in grouped[day]:
            by_plan[str(row.get("plan_type") or "").strip()].append(row)
        if set(by_plan)!=set(amc_expenses.CANARA_PLANS):continue
        if any(len(by_plan[x])!=1 for x in by_plan):
            raise ValueError("Canara exact Mid Cap plan rows are duplicated")
        plans={}
        for published,plan in amc_expenses.CANARA_PLANS.items():
            row=by_plan[published][0]
            plans[plan]={
                "base_expense_ratio":amc_expenses._value(row,"base_ter"),
                "brokerage":0.0,
                "transaction_cost":0.0,
                "statutory_levies":0.0,
                "ter":amc_expenses._value(row,"total_ter"),
            }
        # Canara publishes BER and TER directly; components between them are not
        # exposed in this API, so do not synthesize component reconciliation.
        for plan in ("Regular","Direct"):
            if plans[plan]["ter"]+1e-9<plans[plan]["base_expense_ratio"]:
                raise ValueError("Canara exact Mid Cap Total TER is below BER")
        return {
            "family":family,"amc":TARGETS[family],"status":"recovered",
            "as_of":day,"plans":plans,"source":url,"sha256":_sha(body),
            "identity":{"scheme_name":next(iter(identities))[0],"scheme_code":next(iter(identities))[1]},
            "evidence_type":"official_amc_json",
        }
    raise ValueError("Canara exact Mid Cap API has no complete plan pair")


def _hsbc(family,today):
    content=amc_expenses._hsbc_workbook()
    book=openpyxl.load_workbook(io.BytesIO(content),data_only=True,read_only=True)
    try:
        if "TER" not in book.sheetnames:raise ValueError("HSBC TER sheet missing")
        rows=book["TER"].iter_rows(values_only=True)
        next(rows,None);title=next(rows,None);header=next(rows,None)
        if not title or str(title[0] or "").strip()!="Total Expense Ratio (TER) for HSBC Mutual Fund":
            raise ValueError("HSBC TER workbook title changed")
        if not header or tuple(str(x or "").strip() for x in header[:14])!=amc_expenses._HSBC_HEADER:
            raise ValueError("HSBC TER workbook columns changed")
        matches=defaultdict(list);identities=set()
        for row in rows:
            if len(row)<14 or not _same_family(row[2],family):continue
            day=amc_expenses._hsbc_day(row[3])
            if day is None or day>today:continue
            identities.add((str(row[0] or "").strip(),str(row[1] or "").strip(),str(row[2] or "").strip()))
            matches[day.isoformat()].append(row)
        if len(identities)!=1:raise ValueError("HSBC exact Mid Cap identity is not unique")
        day,row=_latest_exact(matches,"HSBC TER workbook")
        plans={
            "Regular":{
                "base_expense_ratio":amc_expenses._hsbc_number(row[4],"Regular BER"),
                "brokerage":amc_expenses._hsbc_number(row[5],"Regular brokerage"),
                "transaction_cost":amc_expenses._hsbc_number(row[6],"Regular transaction cost"),
                "statutory_levies":amc_expenses._hsbc_number(row[7],"Regular statutory levies"),
                "ter":amc_expenses._hsbc_number(row[8],"Regular Total TER"),
            },
            "Direct":{
                "base_expense_ratio":amc_expenses._hsbc_number(row[9],"Direct BER"),
                "brokerage":amc_expenses._hsbc_number(row[10],"Direct brokerage"),
                "transaction_cost":amc_expenses._hsbc_number(row[11],"Direct transaction cost"),
                "statutory_levies":amc_expenses._hsbc_number(row[12],"Direct statutory levies"),
                "ter":amc_expenses._hsbc_number(row[13],"Direct Total TER"),
            },
        }
        _reconcile(plans,"HSBC Mid Cap")
        code,nsdl,name=next(iter(identities))
        return {
            "family":family,"amc":TARGETS[family],"status":"recovered","as_of":day,
            "plans":plans,"source":amc_expenses.HSBC_TER_LINK,"sha256":_sha(content),
            "identity":{"scheme_name":name,"scheme_code":code,"nsdl_scheme_code":nsdl},
            "evidence_type":"official_amc_linked_workbook",
        }
    finally:
        book.close()


def _icici_fallback_sources(today):
    """Current then previous month, on ICICI's reviewed public TER workbook path."""
    fy=amc_expenses._icici_financial_year(today)
    month=date(today.year,today.month,1)
    previous=(month-timedelta(days=1)).replace(day=1)
    for period in (month,previous):
        # A fiscal-year rollover can make the previous month belong to another folder.
        folder=amc_expenses._icici_financial_year(period)
        yield (
            amc_expenses.ICICI_FILE_BASE
            + "/financials-disclosures-files/Files/Total%20Expense%20Ratio/"
            + folder
            + "/TotalExpenseRatio"
            + period.strftime("%b%Y")
            + ".xlsx"
        )


def _icici(family,today):
    source=None;content=None;api_error=None
    try:
        categories=amc_expenses._icici_api(amc_expenses.ICICI_CATEGORIES_API+"?userType=Investor")
        parent,child,fy=amc_expenses._icici_category_ids(categories,today)
        payload={
            "categoryId":child,"userType":"Investor","fileType":"All","page":"1","size":"20",
            "filter":[{"SHOW":[amc_expenses.ICICI_TER_SHOW]},{"FINANCIAL_YEAR":[fy]}],"search":"",
        }
        data=amc_expenses._icici_api(amc_expenses.ICICI_FILES_API,body=payload)
        if not isinstance(data,dict):raise ValueError("ICICI TER file list changed format")
        source=amc_expenses._icici_select_file(data.get("files"),parent,child,fy,today)
        content,_,_=providers.fetch(source,archive=False,max_bytes=5*1024*1024)
        if not content.startswith(b"PK"):raise ValueError("ICICI TER source is not XLSX")
    except Exception as exc:
        api_error=(str(exc) or type(exc).__name__)[:200]
        rejected_candidates=[];first_rejected=None
        for candidate in _icici_fallback_sources(today):
            body=None
            try:
                body,_,_=providers.fetch(candidate,archive=False,max_bytes=5*1024*1024)
                if not body.startswith(b"PK"):
                    raise ValueError("ICICI TER source is not XLSX")
                # Validate each candidate before selecting it so a rejected
                # current-month placeholder cannot hide the previous workbook.
                result=_icici_workbook_result(family,today,candidate,body,api_error)
            except Exception as candidate_exc:
                rejection={"source":candidate,"error":(str(candidate_exc) or type(candidate_exc).__name__)[:300]}
                if body is not None:
                    rejection.update({"sha256":_sha(body),"bytes":len(body)})
                    if first_rejected is None:
                        first_rejected=(candidate_exc,candidate,body)
                rejected_candidates.append(rejection)
                continue
            if rejected_candidates:
                result["rejected_candidates"]=rejected_candidates
            return result
        if first_rejected is not None:
            first_exc,source,content=first_rejected
            raise _IciciWorkbookError(
                str(first_exc) or type(first_exc).__name__,source,content,
                api_error=api_error,rejected_candidates=rejected_candidates,
            ) from first_exc
        raise _IciciWorkbookError(
            "ICICI financial disclosure API and reviewed TER workbook paths are unavailable",
            api_error=api_error,rejected_candidates=rejected_candidates,
        ) from exc
    try:
        return _icici_workbook_result(family,today,source,content,api_error)
    except Exception as exc:
        raise _IciciWorkbookError(str(exc) or type(exc).__name__,source,content) from exc


def _icici_workbook_result(family,today,source,content,api_error):
    rows=amc_expenses._icici_inline_strings(content)
    if len(rows)<4:raise ValueError("ICICI TER workbook has no usable rows")
    for col,expected in amc_expenses._ICICI_HEADER.items():
        if rows[2].get(col,"").strip()!=expected:
            raise ValueError("ICICI TER workbook columns changed")
    matches=defaultdict(list);names=set()
    for row in rows[3:]:
        published=row.get("A","").strip()
        if not _same_family(published,family):continue
        try:d=datetime.strptime(row.get("B","").strip(),"%d/%m/%Y").date()
        except ValueError:continue
        if d<=today:
            names.add(published);matches[d.isoformat()].append(row)
    if len(names)!=1:raise ValueError("ICICI exact Mid Cap scheme name is not unique")
    day,row=_latest_exact(matches,"ICICI TER workbook")
    plans={
        "Regular":{
            "base_expense_ratio":amc_expenses._icici_percent(row.get("C"),"Regular BER"),
            "brokerage":amc_expenses._icici_percent(row.get("D"),"Regular brokerage"),
            "transaction_cost":amc_expenses._icici_percent(row.get("E"),"Regular transaction cost"),
            "statutory_levies":amc_expenses._icici_percent(row.get("F"),"Regular statutory levies"),
            "ter":amc_expenses._icici_percent(row.get("G"),"Regular Total TER"),
        },
        "Direct":{
            "base_expense_ratio":amc_expenses._icici_percent(row.get("H"),"Direct BER"),
            "brokerage":amc_expenses._icici_percent(row.get("I"),"Direct brokerage"),
            "transaction_cost":amc_expenses._icici_percent(row.get("J"),"Direct transaction cost"),
            "statutory_levies":amc_expenses._icici_percent(row.get("K"),"Direct statutory levies"),
            "ter":amc_expenses._icici_percent(row.get("L"),"Direct Total TER"),
        },
    }
    _reconcile(plans,"ICICI Mid Cap")
    return {
        "family":family,"amc":TARGETS[family],"status":"recovered","as_of":day,
        "plans":plans,"source":source,"sha256":_sha(content),
        "identity":{"scheme_name":next(iter(names))},
        "evidence_type":"official_amc_workbook",
        "discovery_channel":"financial_disclosure_api" if api_error is None else "reviewed_monthly_workbook_fallback",
        "discovery_api_error":api_error,
    }


def _invesco(family,today):
    raw,_,_=providers.fetch(amc_expenses.INVESCO_PLANS_API,archive=False,max_bytes=1024*1024)
    plans_list=json.loads(raw)
    published=next((x for x in plans_list if _same_family(x,family)),None) if isinstance(plans_list,list) else None
    if published is None:raise ValueError("Invesco exact Mid Cap family absent from public TER selector")
    selected=None;records=None
    for period in amc_expenses._invesco_periods(today):
        url=amc_expenses.INVESCO_TER_API+"?"+urlencode({
            "title":published,"fincialYear":amc_expenses._invesco_financial_year_start(period),
            "month":period.month,
        })
        body,_,_=providers.fetch(url,archive=False,max_bytes=2*1024*1024)
        candidate=json.loads(body)
        if isinstance(candidate,list) and candidate:
            selected=(url,body);records=candidate;break
    if selected is None:raise ValueError("Invesco exact Mid Cap API returned no current/previous-month rows")
    matches=defaultdict(list);ids=set()
    for row in records:
        if not isinstance(row,dict) or not _same_family(row.get("Scheme Name",""),family):continue
        try:d=datetime.strptime(str(row.get("TER Date(DD/MM/YYYY)") or "").strip(),"%d/%m/%Y").date()
        except ValueError:continue
        if d<=today:
            ids.add((str(row.get("Scheme Name") or "").strip(),str(row.get("NSDL Scheme Code") or "").strip()))
            matches[d.isoformat()].append(row)
    if len(ids)!=1:raise ValueError("Invesco exact Mid Cap identity is not unique")
    day,row=_latest_exact(matches,"Invesco TER API")
    out={}
    for plan,fields in amc_expenses._INVESCO_FIELDS.items():
        out[plan]={metric:amc_expenses._invesco_percent(row.get(field),f"{plan} {metric}") for metric,field in fields.items()}
    _reconcile(out,"Invesco Mid Cap")
    name,nsdl=next(iter(ids))
    return {
        "family":family,"amc":TARGETS[family],"status":"recovered","as_of":day,
        "plans":out,"source":selected[0],"sha256":_sha(selected[1]),
        "identity":{"scheme_name":name,"nsdl_scheme_code":nsdl},
        "evidence_type":"official_amc_json",
    }


def _jm(family,today):
    raw,_,_=providers.fetch(amc_expenses.JM_TER_API,body=amc_expenses.JM_TER_REQUEST,
                            archive=False,max_bytes=4*1024*1024)
    records=jm_portfolios._decrypt(raw)
    matches=defaultdict(list);ids=set()
    for row in records:
        if not isinstance(row,dict) or not _same_family(row.get("Scheme",""),family):continue
        try:d=datetime.fromisoformat(str(row.get("TERDate") or "").strip().replace("Z","+00:00")).date()
        except ValueError:continue
        if d<=today:
            ids.add((str(row.get("Scheme") or "").strip(),str(row.get("Schemecode") or "").strip(),str(row.get("NsdlSchemeCode") or "").strip()))
            matches[d.isoformat()].append(row)
    if len(ids)!=1:raise ValueError("JM exact Mid Cap identity is not unique")
    day,row=_latest_exact(matches,"JM TER API")
    out={}
    for plan,fields in amc_expenses._JM_FIELDS.items():
        out[plan]={metric:amc_expenses._jm_percent(row.get(field),f"{plan} {metric}") for metric,field in fields.items()}
    _reconcile(out,"JM Mid Cap")
    name,code,nsdl=next(iter(ids))
    evidence=json.dumps(records,sort_keys=True,separators=(",",":"),ensure_ascii=False).encode()
    return {
        "family":family,"amc":TARGETS[family],"status":"recovered","as_of":day,
        "plans":out,"source":amc_expenses.JM_TER_API,"sha256":_sha(evidence),
        "identity":{"scheme_name":name,"scheme_code":code,"nsdl_scheme_code":nsdl},
        "evidence_type":"official_amc_json",
    }


def _mahindra(family,today):
    raw,_,_=providers.fetch(amc_expenses.MAHINDRA_DOWNLOADS_API,archive=False,max_bytes=4*1024*1024)
    tree=amc_expenses._mahindra_decrypt_downloads(raw)
    source=amc_expenses._mahindra_select_ter_file(tree,today)
    content,_,_=providers.fetch(source,archive=False,max_bytes=5*1024*1024)
    book=openpyxl.load_workbook(io.BytesIO(content),data_only=True,read_only=True)
    try:
        if book.sheetnames!=["Sheet1"]:raise ValueError("Mahindra TER workbook layout changed")
        rows=book["Sheet1"].iter_rows(values_only=True)
        h1=next(rows,None);h2=next(rows,None)
        if tuple("" if v is None else str(v).strip() for v in (h1 or ())[:13])!=amc_expenses._MAHINDRA_HEADER_1:
            raise ValueError("Mahindra TER workbook top header changed")
        if tuple("" if v is None else str(v).strip() for v in (h2 or ())[:13])!=amc_expenses._MAHINDRA_HEADER_2:
            raise ValueError("Mahindra TER workbook metric columns changed")
        matches=defaultdict(list);ids=set()
        for row in rows:
            if len(row)<13 or not _same_family(row[1],family):continue
            try:d=datetime.strptime(str(row[2] or "").strip(),"%d-%b-%Y").date()
            except ValueError:continue
            if d<=today:
                ids.add((str(row[0] or "").strip(),str(row[1] or "").strip()))
                matches[d.isoformat()].append(row)
        if len(ids)!=1:raise ValueError("Mahindra exact Mid Cap identity is not unique")
        day,row=_latest_exact(matches,"Mahindra TER workbook")
        out={
            "Regular":{
                "base_expense_ratio":amc_expenses._mahindra_percent(row[3],"Regular BER"),
                "brokerage":amc_expenses._mahindra_percent(row[4],"Regular brokerage"),
                "transaction_cost":amc_expenses._mahindra_percent(row[5],"Regular transaction cost"),
                "statutory_levies":amc_expenses._mahindra_percent(row[6],"Regular statutory levies"),
                "ter":amc_expenses._mahindra_percent(row[7],"Regular Total TER"),
            },
            "Direct":{
                "base_expense_ratio":amc_expenses._mahindra_percent(row[8],"Direct BER"),
                "brokerage":amc_expenses._mahindra_percent(row[9],"Direct brokerage"),
                "transaction_cost":amc_expenses._mahindra_percent(row[10],"Direct transaction cost"),
                "statutory_levies":amc_expenses._mahindra_percent(row[11],"Direct statutory levies"),
                "ter":amc_expenses._mahindra_percent(row[12],"Direct Total TER"),
            },
        }
        _reconcile(out,"Mahindra Mid Cap")
        nsdl,name=next(iter(ids))
        return {
            "family":family,"amc":TARGETS[family],"status":"recovered","as_of":day,
            "plans":out,"source":source,"sha256":_sha(content),
            "identity":{"scheme_name":name,"nsdl_scheme_code":nsdl},
            "evidence_type":"official_amc_workbook",
        }
    finally:
        book.close()


def _mirae(family,today):
    start=today-timedelta(days=35)
    req={"request":{"modulename":"TotalExpenseRatio","title":"","fromdate":start.isoformat(),
                    "todate":today.isoformat(),"pgno":1,"pgsize":100}}
    raw,_,_=providers.fetch(amc_expenses.MIRAE_TER_API,body=req,archive=False,max_bytes=4*1024*1024)
    payload=json.loads(raw)
    _,source=amc_expenses._mirae_select_download(payload,today)
    content,_,_=providers.fetch(source,archive=False,max_bytes=5*1024*1024)
    book=xlrd.open_workbook(file_contents=content)
    try:
        if book.sheet_names()!=["Report"]:raise ValueError("Mirae TER workbook layout changed")
        sheet=book.sheet_by_name("Report")
        matches=defaultdict(list);ids=set()
        for i in range(2,sheet.nrows):
            row=sheet.row_values(i,0,13)
            if len(row)<13 or not _same_family(row[1],family):continue
            try:
                d=xlrd.xldate_as_datetime(row[2],book.datemode).date() if isinstance(row[2],(int,float)) else datetime.strptime(str(row[2]).strip(),"%d/%m/%Y").date()
            except (ValueError,TypeError,xlrd.XLDateError):continue
            if d<=today:
                ids.add((str(row[0] or "").strip(),str(row[1] or "").strip()))
                matches[d.isoformat()].append(row)
        if len(ids)!=1:raise ValueError("Mirae exact Mid Cap identity is not unique")
        day,row=_latest_exact(matches,"Mirae TER workbook")
        out={
            "Regular":{
                "base_expense_ratio":amc_expenses._mirae_percent(row[3],"Regular BER"),
                "brokerage":amc_expenses._mirae_percent(row[4],"Regular brokerage"),
                "transaction_cost":amc_expenses._mirae_percent(row[5],"Regular transaction cost"),
                "statutory_levies":amc_expenses._mirae_percent(row[6],"Regular statutory levies"),
                "ter":amc_expenses._mirae_percent(row[7],"Regular Total TER"),
            },
            "Direct":{
                "base_expense_ratio":amc_expenses._mirae_percent(row[8],"Direct BER"),
                "brokerage":amc_expenses._mirae_percent(row[9],"Direct brokerage"),
                "transaction_cost":amc_expenses._mirae_percent(row[10],"Direct transaction cost"),
                "statutory_levies":amc_expenses._mirae_percent(row[11],"Direct statutory levies"),
                "ter":amc_expenses._mirae_percent(row[12],"Direct Total TER"),
            },
        }
        _reconcile(out,"Mirae Mid Cap")
        nsdl,name=next(iter(ids))
        return {
            "family":family,"amc":TARGETS[family],"status":"recovered","as_of":day,
            "plans":out,"source":source,"sha256":_sha(content),
            "identity":{"scheme_name":name,"nsdl_scheme_code":nsdl},
            "evidence_type":"official_amc_workbook",
        }
    finally:
        book.release_resources()


COLLECTORS={
    "BANK OF INDIA MID CAP FUND":_boi,
    "Canara Robeco Mid Cap Fund":_canara,
    "HSBC Midcap Fund":_hsbc,
    "ICICI Prudential Mid Cap Fund":_icici,
    "Invesco India Mid Cap Fund":_invesco,
    "JM Mid Cap Fund":_jm,
    "Mahindra Manulife Mid Cap Fund":_mahindra,
    "Mirae Asset Midcap Fund":_mirae,
}


def collect(today=None):
    today=today or date.today()
    staged={row["family"] for row in db.rows(
        "SELECT DISTINCT family FROM category_staged_schemes WHERE category='mid-cap'"
    )}
    results=[];errors=[]
    for family,collector in COLLECTORS.items():
        if family not in staged:
            errors.append({"family":family,"error":"staged family identity missing"})
            continue
        try:
            result=collector(family,today)
            result["direct_ter"]=result["plans"]["Direct"]["ter"]
            result["regular_ter"]=result["plans"]["Regular"]["ter"]
            results.append(result)
        except Exception as exc:
            error={"family":family,"amc":TARGETS[family],
                   "error":(str(exc) or type(exc).__name__)[:300]}
            if isinstance(exc,(_IciciWorkbookError,_BoiWorkbookError)):
                error.update(exc.source_evidence)
            errors.append(error)
    return {
        "built_at":db.now(),
        "staged_category":"mid-cap",
        "targets":len(COLLECTORS),
        "recovered":len(results),
        "failed":len(errors),
        "results":results,
        "errors":errors,
        "production_writes":0,
        "public_export_enabled":False,
        "notes":[
            "All source requests are read-only and archive=False; this audit does not insert metrics, documents or fetch records.",
            "A result requires the first-party source to contain the staged Mid Cap family identity and a complete Regular/Direct TER pair.",
            "No Small Cap TER value, NSDL code or scheme code is reused for a Mid Cap fund.",
        ],
    }
