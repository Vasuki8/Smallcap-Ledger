"""Read-only first-party TER recovery batch 2 for staged Mid Cap funds."""
from __future__ import annotations

from collections import defaultdict
from datetime import date,datetime
import hashlib,io,re

import openpyxl,xlrd

from . import amfi_metrics,db,providers


HELIOS_FAMILY="Helios Mid Cap Fund"
HELIOS_URL="https://www.heliosmf.in/wp-content/uploads/2026/09/TER_sep_2026_vsvefi.xls"
KOTAK_FAMILY="Kotak Mid Cap Fund"
KOTAK_URL="https://vatseelabs-s3.kotakmf.com/TER/TER-2026-2027.xlsx"
WEALTH_FAMILY="The Wealth Company Mid Cap Fund"
WEALTH_URL="https://www.wealthcompanyamc.in/api/ter-values/year/2026/export"


def _same_family(a,b):
    return amfi_metrics.normalized(a)==amfi_metrics.normalized(b)


def _num(value,label):
    if value is None or str(value).strip() in ("","-","NA","N/A"):
        return 0.0
    parsed=float(value)
    if not 0<=parsed<=10:
        raise ValueError(f"{label} is outside accepted range")
    return parsed


def _validate(plans,label):
    for plan in ("Regular","Direct"):
        values=plans[plan]
        if values["ter"]+1e-9<values["base_expense_ratio"]:
            raise ValueError(f"{label} {plan} TER below BER")
        total=sum(values[k] for k in ("base_expense_ratio","brokerage","transaction_cost","statutory_levies"))
        if abs(total-values["ter"])>0.03:
            raise ValueError(f"{label} {plan} components do not reconcile")
    return plans


def _latest(rows,label):
    if not rows:raise ValueError(f"{label} contains no exact Mid Cap rows")
    day=max(rows)
    if len(rows[day])!=1:raise ValueError(f"{label} duplicate exact rows for {day}")
    return day,rows[day][0]


def helios(today=None):
    today=today or date.today()
    body,_,_=providers.fetch(HELIOS_URL,archive=False,max_bytes=8*1024*1024)
    book=xlrd.open_workbook(file_contents=body)
    try:
        if "Mid Cap" not in book.sheet_names():raise ValueError("Helios Mid Cap sheet missing")
        sh=book.sheet_by_name("Mid Cap")
        matches=defaultdict(list);ids=set()
        for i in range(5,sh.nrows):
            row=sh.row_values(i,0,13)
            if len(row)<13 or not _same_family(row[1],HELIOS_FAMILY):continue
            try:d=xlrd.xldate_as_datetime(row[2],book.datemode).date()
            except Exception:continue
            if d<=today:
                ids.add((str(row[0]).strip(),str(row[1]).strip()))
                matches[d.isoformat()].append(row)
        if len(ids)!=1:raise ValueError("Helios exact identity not unique")
        day,row=_latest(matches,"Helios TER")
        plans={
          "Regular":{"base_expense_ratio":_num(row[3],"BER"),"brokerage":_num(row[4],"brokerage"),"transaction_cost":_num(row[5],"txn"),"statutory_levies":_num(row[6],"levies"),"ter":_num(row[7],"TER")},
          "Direct":{"base_expense_ratio":_num(row[8],"BER"),"brokerage":_num(row[9],"brokerage"),"transaction_cost":_num(row[10],"txn"),"statutory_levies":_num(row[11],"levies"),"ter":_num(row[12],"TER")},
        }
        _validate(plans,"Helios")
        nsdl,name=next(iter(ids))
        return {"family":HELIOS_FAMILY,"status":"recovered","as_of":day,"plans":plans,"direct_ter":plans["Direct"]["ter"],"regular_ter":plans["Regular"]["ter"],"source":HELIOS_URL,"sha256":hashlib.sha256(body).hexdigest(),"identity":{"scheme_name":name,"nsdl_scheme_code":nsdl}}
    finally:book.release_resources()


def kotak(today=None):
    today=today or date.today()
    body,_,_=providers.fetch(KOTAK_URL,archive=False,max_bytes=20*1024*1024)
    book=openpyxl.load_workbook(io.BytesIO(body),data_only=True,read_only=True)
    try:
        matches=defaultdict(list);ids=set()
        for name in book.sheetnames:
            sh=book[name]
            for row in sh.iter_rows(min_row=3,values_only=True):
                if len(row)<13 or not _same_family(row[0],KOTAK_FAMILY):continue
                try:d=datetime.strptime(str(row[1]).strip(),"%d/%m/%Y").date()
                except ValueError:continue
                if d<=today:
                    ids.add((str(row[0]).strip(),str(row[12]).strip()))
                    matches[d.isoformat()].append(row)
        if len(ids)!=1:raise ValueError("Kotak exact identity not unique")
        day,row=_latest(matches,"Kotak TER")
        plans={
          "Regular":{"base_expense_ratio":_num(row[2],"BER"),"brokerage":_num(row[3],"brokerage"),"transaction_cost":_num(row[4],"txn"),"statutory_levies":_num(row[5],"levies"),"ter":_num(row[6],"TER")},
          "Direct":{"base_expense_ratio":_num(row[7],"BER"),"brokerage":_num(row[8],"brokerage"),"transaction_cost":_num(row[9],"txn"),"statutory_levies":_num(row[10],"levies"),"ter":_num(row[11],"TER")},
        }
        _validate(plans,"Kotak")
        name,nsdl=next(iter(ids))
        return {"family":KOTAK_FAMILY,"status":"recovered","as_of":day,"plans":plans,"direct_ter":plans["Direct"]["ter"],"regular_ter":plans["Regular"]["ter"],"source":KOTAK_URL,"sha256":hashlib.sha256(body).hexdigest(),"identity":{"scheme_name":name,"nsdl_scheme_code":nsdl}}
    finally:book.close()


def wealth(today=None):
    today=today or date.today()
    body,_,_=providers.fetch(WEALTH_URL,archive=False,max_bytes=20*1024*1024)
    book=openpyxl.load_workbook(io.BytesIO(body),data_only=True,read_only=True)
    try:
        if len(book.sheetnames)!=1:raise ValueError("Wealth TER workbook sheet count changed")
        sh=book[book.sheetnames[0]]
        matches=defaultdict(list);ids=set()
        for row in sh.iter_rows(min_row=2,values_only=True):
            if len(row)<13 or not _same_family(row[1],WEALTH_FAMILY):continue
            raw=row[2]
            try:
                d=raw.date() if isinstance(raw,datetime) else datetime.fromisoformat(str(raw)).date()
            except Exception:continue
            if d<=today:
                ids.add((str(row[0]).strip(),str(row[1]).strip()))
                matches[d.isoformat()].append(row)
        if len(ids)!=1:raise ValueError("Wealth exact identity not unique")
        day,row=_latest(matches,"Wealth TER")
        plans={
          "Regular":{"base_expense_ratio":_num(row[3],"BER"),"brokerage":_num(row[4],"brokerage"),"transaction_cost":_num(row[5],"txn"),"statutory_levies":_num(row[6],"levies"),"ter":_num(row[7],"TER")},
          "Direct":{"base_expense_ratio":_num(row[8],"BER"),"brokerage":_num(row[9],"brokerage"),"transaction_cost":_num(row[10],"txn"),"statutory_levies":_num(row[11],"levies"),"ter":_num(row[12],"TER")},
        }
        _validate(plans,"Wealth")
        nsdl,name=next(iter(ids))
        return {"family":WEALTH_FAMILY,"status":"recovered","as_of":day,"plans":plans,"direct_ter":plans["Direct"]["ter"],"regular_ter":plans["Regular"]["ter"],"source":WEALTH_URL,"sha256":hashlib.sha256(body).hexdigest(),"identity":{"scheme_name":name,"nsdl_scheme_code":nsdl}}
    finally:book.close()


def collect(today=None):
    today=today or date.today()
    staged={r["family"] for r in db.rows("SELECT DISTINCT family FROM category_staged_schemes WHERE category='mid-cap'")}
    results=[];errors=[]
    for family,fn in ((HELIOS_FAMILY,helios),(KOTAK_FAMILY,kotak),(WEALTH_FAMILY,wealth)):
        if family not in staged:
            errors.append({"family":family,"error":"staged identity missing"});continue
        try:results.append(fn(today))
        except Exception as exc:errors.append({"family":family,"error":(str(exc) or type(exc).__name__)[:300]})
    return {"built_at":db.now(),"staged_category":"mid-cap","targets":3,"recovered":len(results),"failed":len(errors),"results":results,"errors":errors,"production_writes":0,"public_export_enabled":False}
