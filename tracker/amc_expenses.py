"""Official AMC expense disclosures that are not currently present in AMFI's Small Cap feed."""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta
import base64
import hashlib
import io
import json
import re
import shutil
import subprocess
import uuid
import zipfile
import xml.etree.ElementTree as ET
from urllib.parse import quote, unquote, urlencode, urljoin, urlparse

import httpx
import openpyxl
import xlrd
from bs4 import BeautifulSoup
from pypdf import PdfReader

from . import db, jm_portfolios
from .providers import fetch, number, public_url

CANARA_FAMILY = "Canara Robeco Small Cap Fund"
CANARA_SCHEME_CODE = "SC"
CANARA_PAGE = "https://www.canararobeco.com/expense-ratio"
CANARA_API = "https://www.canararobeco.com/wp-json/ter/v1/records"
CANARA_PLANS = {"Regular Plan": "Regular", "Direct Plan": "Direct"}

GROWW_FAMILY = "Groww Small Cap Fund"
GROWW_PUBLISHED_NAME = "Groww Smallcap Fund"
GROWW_BER_PAGE = "https://www.growwmf.in/downloads/expense-ratio"
GROWW_BER_HOST = "assets-netstorage.growwmf.in"
_GROWW_BER_TITLE = re.compile(r"^(\d+)\.\s*Notice\s*-\s*Change\s+in\s+BER\.pdf$", re.I)

ICICI_FAMILY = "ICICI Prudential Small Cap Fund"
ICICI_TER_PAGE = "https://www.icicipruamc.com/about-us/financials-&-disclosures?currentTabFilter=Total%20Expense%20Ratio"
ICICI_API_BASE = "https://apimf.icicipruamc.com"
ICICI_CATEGORIES_API = ICICI_API_BASE + "/fds/v1/categories"
ICICI_FILES_API = ICICI_API_BASE + "/fds/v1/files"
ICICI_FILE_BASE = "https://app.beta.icicipruamc.com/blob"
ICICI_TER_CATEGORY_CODE = "TOTAL_EXPENSE_RATIO"
ICICI_TER_SUBCATEGORY_CODE = "TOTAL_EXPENSE_RATIO"
ICICI_TER_TITLE = "Total Expense Ratio"
ICICI_TER_SHOW = "TER Details"
_ICICI_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

INVESCO_FAMILY = "Invesco India Small Cap Fund"
INVESCO_NSDL_CODE = "INVM/O/E/SCF/18/07/0030"
INVESCO_TER_PAGE = "https://www.invescomutualfund.com/statutory-disclosures/ter-mutual-fund-since-2026/ter"
INVESCO_PLANS_API = "https://www.invescomutualfund.com/api/Common/GetAllPlans"
INVESCO_TER_API = "https://www.invescomutualfund.com/api/TotalExpenseRatioOfMutualFundSchemePolicy/GetTERExpenseData"

JM_FAMILY = "Jm Small Cap Fund"
JM_PUBLISHED_NAME = "JM Small Cap Fund"
JM_SCHEME_CODE = "SC"
JM_NSDL_CODE = "JMFI/O/E/SCF/23/11/0016"
JM_TER_PAGE = "https://www.jmfinancialmf.com/Scheme-Expense-Ratio"
JM_TER_API = jm_portfolios.API_BASE + "GetTerPageLatest"
JM_TER_REQUEST = {"IICategory": 0, "IVFundCode": ""}

MIRAE_FAMILY = "Mirae Asset Small Cap Fund"
MIRAE_NSDL_CODE = "MIRA/O/E/SCF/24/10/0075"
MIRAE_TER_PAGE = "https://www.miraeassetmf.co.in/downloads/statutory-disclosure/total-expense-ratio"
MIRAE_TER_API = "https://www.miraeassetmf.co.in/AjaxService/GetDownloadsData"
MIRAE_TER_BASE = "https://www.miraeassetmf.co.in"
_MIRAE_FILE = re.compile(r"^(?:R_)?IN_MF_EXPENSE_RATIO_SEBI_V3_(\d{8})\.xls$", re.I)
_MIRAE_TITLE = re.compile(r"^Total Expense Ratio -(\d{2} [A-Za-z]{3} 20\d{2})$")
_MIRAE_HEADER_1 = (
    "NSDL Scheme Code",
    "Scheme Name",
    "TER Date (DD/MM/ YYYY)",
    "Regular",
    "",
    "",
    "",
    "",
    "Direct",
    "",
    "",
    "",
    "",
)
_MIRAE_HEADER_2 = (
    "",
    "",
    "",
    "Regular Plan - Base Expense Ratio (BER) (%)",
    "Regular Plan - Brokerage cost (%)",
    "Regular Plan - Transaction Cost incurred for the purpose of execution of trade (%)",
    "Regular Plan - Statutory Levies (including GST) (%)",
    "Regular Plan - Total TER (%)",
    "Direct Plan - Base Expense Ratio (BER) (%)",
    "Direct Plan - Brokerage cost (%)",
    "Direct Plan - Transaction Cost incurred for the purpose of execution of trade (%)",
    "Direct Plan - Statutory Levies (including GST) (%)",
    "Direct Plan - Total TER (%)",
)

MAHINDRA_FAMILY = "Mahindra Manulife Small Cap Fund"
MAHINDRA_NSDL_CODE = "MAHM/O/E/SCF/22/07/0020"
MAHINDRA_DOWNLOADS_PAGE = "https://www.mahindramanulife.com/downloads"
MAHINDRA_DOWNLOADS_API = "https://investorapi.mahindramanulife.com/api/v1/web/preLogin/downloads"
_MAHINDRA_AES_KEY = b"mahindra2024mahindra2024mahindra"
_MAHINDRA_AES_IV = b"hasnainsheikh202"
_MAHINDRA_TOP_CATEGORY = "MANDATORY DISCLOSURES"
_MAHINDRA_TER_CATEGORY = "Total Expense Ratio of Mutual Fund Schemes"
_MAHINDRA_TER_SUBCATEGORY = "Total Expense Ratio"
_MAHINDRA_HEADER_1 = (
    "NSDL Scheme Code",
    "Name of Scheme",
    "Date (DD/MM/YYYY)",
    "Regular",
    "",
    "",
    "",
    "",
    "Direct",
    "",
    "",
    "",
    "",
)
_MAHINDRA_HEADER_2 = (
    "",
    "",
    "",
    "Base Expense Ratio (BER) (%)1",
    "Brokerage cost (%)2",
    "Transaction Cost incurred for the purpose of execution of trade (%)3",
    "Statutory Levies (including GST) (%)4",
    "Total TER (%)",
    "Base Expense Ratio (BER) (%)1",
    "Brokerage cost (%)2",
    "Transaction Cost incurred for the purpose of execution of trade (%)3",
    "Statutory Levies (including GST) (%)4",
    "Total TER (%)",
)
_JM_FIELDS = {
    "Regular": {
        "base_expense_ratio": "RegularBER",
        "brokerage": "RegularBrokCost",
        "transaction_cost": "RegularTransCost",
        "statutory_levies": "RegularStatLevGST",
        "ter": "RegularTotalTER",
    },
    "Direct": {
        "base_expense_ratio": "DirectBER",
        "brokerage": "DirectBrokCost",
        "transaction_cost": "DirectTransCost",
        "statutory_levies": "DirectStatLevGST",
        "ter": "DirectTotalTER",
    },
}
_INVESCO_FIELDS = {
    "Regular": {
        "base_expense_ratio": "Regular Plan - Base Expense Ratio (BER) (%)",
        "brokerage": "Regular Plan - Brokerage cost (%)",
        "transaction_cost": "Regular Plan - Transaction Cost incurred for the purpose of execution of trade (%)",
        "statutory_levies": "Regular Plan - Statutory Levies (including GST) (%)",
        "ter": "Regular Plan - Total TER (%)",
    },
    "Direct": {
        "base_expense_ratio": "Direct Plan - Base Expense Ratio (BER) (%)",
        "brokerage": "Direct Plan - Brokerage cost (%)",
        "transaction_cost": "Direct Plan - Transaction Cost incurred for the purpose of execution of trade (%)",
        "statutory_levies": "Direct Plan - Statutory Levies (including GST) (%)",
        "ter": "Direct Plan - Total TER (%)",
    },
}
_ICICI_TITLE = re.compile(r"^TotalExpenseRatio([A-Za-z]+)(20\d{2})$")
_ICICI_FILE = re.compile(r"^TotalExpenseRatio([A-Za-z]+)(20\d{2})\.xlsx$", re.I)
_ICICI_HEADER = {
    "A": "Scheme Name",
    "B": "Date (DD/MM/YYYY)",
    "C": "Base Expense Ratio (BER) (%)",
    "D": "Brokerage cost (%)",
    "E": "Transaction Cost incurred for the purpose of execution of trade (%)",
    "F": "Statutory Levies (including GST) (%)",
    "G": "Total TER (%)",
    "H": "Base Expense Ratio (BER) (%)",
    "I": "Brokerage cost (%)",
    "J": "Transaction Cost incurred for the purpose of execution of trade (%)",
    "K": "Statutory Levies (including GST) (%)",
    "L": "Total TER (%)",
}

HSBC_FAMILY = "HSBC Small Cap Fund"
HSBC_SCHEME_CODE = "HEMIDF"
HSBC_NSDL_CODE = "LTMF/O/E/SCF/14/02/0023"
HSBC_TER_LINK = "https://digital.camsonline.com/dnlresult/hsbc_ter_report.xlsx"
HSBC_CAMS_API = "https://digital.camsonline.com/api/v1/camsonline"
_HSBC_IV = b"globalaesvectors"
_HSBC_REQUEST_SECRET = "TkVJTEhobWFj"
_HSBC_RESPONSE_SECRET = "UkRYTElobWFj"
_HSBC_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_HSBC_HEADER = (
    "Scheme Code",
    "NSDL Scheme Code",
    "Scheme Name",
    "TER Date",
    "Regular Plan - Base Expense Ratio (BER) (%)",
    "Regular Plan - Brokerage cost (%)",
    "Regular Plan - Transaction Cost incurred for the purpose of execution of trade (%)",
    "Regular Plan - Statutory Levies (including GST) (%)",
    "Regular Plan - Total TER (%)",
    "Direct Plan - Base Expense Ratio (BER) (%)",
    "Direct Plan - Brokerage cost (%)",
    "Direct Plan - Transaction Cost incurred for the purpose of execution of trade (%)",
    "Direct Plan - Statutory Levies (including GST) (%)",
    "Direct Plan - Total TER (%)",
)


def _value(row, field):
    raw = row.get(field)
    if raw is None or str(raw).strip() in ("", "-", "NA", "N/A"):
        raise ValueError(f"Canara expense row is missing {field}")
    value = number(raw)
    if not 0 <= value <= 5:
        raise ValueError(f"Canara expense {field} is outside the accepted range")
    return value


def parse_canara_records(records, today=None):
    """Return the latest exact two-plan Small Cap BER/TER disclosure.

    The public Canara page uses this first-party JSON API in the browser.
    We retain only the AMC-published base_ter and total_ter fields; components
    are not recomputed into a TER.
    """
    today = today or date.today()
    if not isinstance(records, list):
        raise ValueError("Canara TER API response changed format")

    grouped = defaultdict(list)
    for row in records:
        if not isinstance(row, dict):
            continue
        if str(row.get("scheme_name", "")).strip() != CANARA_FAMILY:
            continue
        if str(row.get("sch_code", "")).strip().upper() != CANARA_SCHEME_CODE:
            continue
        plan = str(row.get("plan_type", "")).strip()
        if plan not in CANARA_PLANS:
            continue
        day = str(row.get("date", "")).strip()
        try:
            parsed = date.fromisoformat(day)
        except ValueError:
            continue
        if parsed > today:
            continue
        grouped[day].append(row)

    for day in sorted(grouped, reverse=True):
        rows = grouped[day]
        by_plan = defaultdict(list)
        for row in rows:
            by_plan[str(row.get("plan_type", "")).strip()].append(row)
        if set(by_plan) != set(CANARA_PLANS):
            continue
        if any(len(by_plan[p]) != 1 for p in CANARA_PLANS):
            raise ValueError(f"Canara TER API returned duplicate plan rows for {day}")

        result = {}
        for published_plan, plan in CANARA_PLANS.items():
            row = by_plan[published_plan][0]
            ber = _value(row, "base_ter")
            ter = _value(row, "total_ter")
            if ter + 1e-9 < ber:
                raise ValueError(f"Canara Total TER is below BER for {published_plan} on {day}")
            result[plan] = {"base_expense_ratio": ber, "ter": ter}
        return day, result

    raise ValueError("Canara TER API returned no complete dated Small Cap plan pair")


def canara(progress=lambda _: None, today=None, lookback_days=7):
    """Collect current Canara Robeco Small Cap BER/TER from its browser API."""
    today = today or date.today()
    if lookback_days < 1:
        raise ValueError("Canara expense lookback must be at least one day")
    if not db.one("SELECT code FROM schemes WHERE family=? LIMIT 1", (CANARA_FAMILY,)):
        return "Canara Small Cap is not in the active universe"

    start = today - timedelta(days=lookback_days - 1)
    url = CANARA_API + "?" + urlencode(
        {"from_date": start.isoformat(), "to_date": today.isoformat()}
    )
    progress("Canara Robeco expense ratios · official TER API")
    body, content_hash, _ = fetch(url, max_bytes=2 * 1024 * 1024)
    try:
        records = json.loads(body)
    except (TypeError, json.JSONDecodeError) as exc:
        raise ValueError("Canara TER API returned invalid JSON") from exc

    day, plans = parse_canara_records(records, today)
    for plan, values in plans.items():
        db.metric(
            CANARA_FAMILY,
            plan,
            "base_expense_ratio",
            day,
            values["base_expense_ratio"],
            "% p.a. · reported by AMC",
            url,
            content_hash,
        )
        db.metric(
            CANARA_FAMILY,
            plan,
            "ter",
            day,
            values["ter"],
            "% p.a. · reported by AMC",
            url,
            content_hash,
        )
    return (
        f"{CANARA_FAMILY}: official AMC BER/TER as of {day} "
        f"(Direct {plans['Direct']['base_expense_ratio']:.2f}%/"
        f"{plans['Direct']['ter']:.2f}% BER/TER)"
    )


def _groww_financial_year(today):
    start = today.year if today.month >= 4 else today.year - 1
    return f"{start} - {start + 1}"


def _groww_ber_links(content, today=None):
    """Return current-financial-year first-party Groww BER notices, newest number first."""
    today = today or date.today()
    soup = BeautifulSoup(content, "html.parser")
    expected_dir = (
        "/compliance_docs/Downloads/Expense Ratio/Notice - Change in TER/"
        + _groww_financial_year(today)
        + "/"
    )
    found = defaultdict(set)
    for link in soup.find_all("a", href=True):
        title = " ".join(link.stripped_strings).strip()
        match = _GROWW_BER_TITLE.fullmatch(title)
        if not match:
            continue
        url = urljoin(GROWW_BER_PAGE, str(link.get("href") or "").strip())
        parsed = urlparse(url)
        path = unquote(parsed.path)
        if (
            parsed.scheme != "https"
            or parsed.hostname != GROWW_BER_HOST
            or not path.startswith(expected_dir)
            or not path.lower().endswith(".pdf")
        ):
            continue
        public_url(url)
        found[int(match.group(1))].add(url)

    if not found:
        raise ValueError("Groww expense page exposes no current-year BER notices")
    if any(len(urls) != 1 for urls in found.values()):
        raise ValueError("Groww expense page has duplicate BER notice identities")
    return [
        (notice, next(iter(found[notice])))
        for notice in sorted(found, reverse=True)
    ]


def _groww_pdf_text(content):
    if not content.startswith(b"%PDF"):
        raise ValueError("Groww BER notice is not a PDF")
    try:
        reader = PdfReader(io.BytesIO(content))
        if len(reader.pages) != 1:
            raise ValueError("Groww BER notice page count changed")
        text = reader.pages[0].extract_text() or ""
    except Exception as exc:
        if isinstance(exc, ValueError):
            raise
        raise ValueError("Groww BER notice could not be parsed") from exc
    if not text.strip():
        raise ValueError("Groww BER notice contains no extractable text")
    return text


def _groww_ber_number(raw, label):
    token = str(raw or "").strip()
    if token.upper() == "NA":
        return None
    try:
        value = number(token)
    except ValueError as exc:
        raise ValueError(f"Groww BER notice has invalid {label}") from exc
    if not 0 <= value <= 5:
        raise ValueError(f"Groww BER notice {label} is outside the accepted range")
    return value


def parse_groww_ber_text(text, today=None):
    """Parse exact Current BER values; deliberately ignore conditional revised BER."""
    today = today or date.today()
    normalized = re.sub(r"\s+", " ", str(text or "")).strip()

    notice = re.search(
        r"Notice\s+no\.\s*(\d+)\s*/\s*(20\d{2})\s*[–-]\s*(20\d{2})",
        normalized,
        re.I,
    )
    if not notice:
        raise ValueError("Groww BER notice identity changed")
    if not re.search(
        r"Change\s+in\s+Base\s+Expense\s+Ratio.*?BER.*?scheme\(s\)\s+of\s+Groww\s+Mutual\s+Fund",
        normalized,
        re.I,
    ):
        raise ValueError("Groww BER notice heading changed")
    if not re.search(
        r"Scheme\(s\)\s+Name\s+Current\s+BER\*\s+Revised\s+BER\*\*\s+"
        r"Direct\s*\(%\)\s+Regular\s*\(%\)\s+Direct\s*\(%\)\s+Regular\s*\(%\)",
        normalized,
        re.I,
    ):
        raise ValueError("Groww BER notice table heading changed")

    as_of_match = re.search(
        r"\*\s*As\s+on\s+([A-Za-z]+\s+\d{1,2},\s*20\d{2})\.",
        normalized,
        re.I,
    )
    signed_match = re.search(
        r"Authorised\s+Signatory\s+Date:\s*([A-Za-z]+\s+\d{1,2},\s*20\d{2})",
        normalized,
        re.I,
    )
    effective_match = re.search(
        r"with\s+effect\s+from\s+([A-Za-z]+\s+\d{1,2},\s*20\d{2})",
        normalized,
        re.I,
    )
    if not as_of_match or not signed_match or not effective_match:
        raise ValueError("Groww BER notice dates changed format")
    try:
        as_of = datetime.strptime(as_of_match.group(1).replace("  ", " "), "%B %d, %Y").date()
        signed = datetime.strptime(signed_match.group(1).replace("  ", " "), "%B %d, %Y").date()
        effective = datetime.strptime(effective_match.group(1).replace("  ", " "), "%B %d, %Y").date()
    except ValueError as exc:
        raise ValueError("Groww BER notice contains an invalid date") from exc
    if as_of > today or signed > today:
        raise ValueError("Groww BER notice contains a future observation/publication date")
    if as_of > signed or effective < signed:
        raise ValueError("Groww BER notice date ordering is invalid")

    row_pattern = re.compile(
        r"Groww\s+Smallcap\s+Fund\s+"
        r"(NA|\d+(?:\.\d+)?)\s+"
        r"(NA|\d+(?:\.\d+)?)\s+"
        r"(NA|\d+(?:\.\d+)?)\s+"
        r"(NA|\d+(?:\.\d+)?)(?:\s*\(No\s+change\))?",
        re.I,
    )
    rows = row_pattern.findall(normalized)
    if len(rows) != 1:
        raise ValueError("Groww BER notice does not contain one exact Smallcap row")
    current_direct, current_regular, revised_direct, revised_regular = rows[0]
    plans = {}
    direct = _groww_ber_number(current_direct, "Direct Current BER")
    regular = _groww_ber_number(current_regular, "Regular Current BER")
    # Validate the published revised columns but never store them as exact current observations.
    _groww_ber_number(revised_direct, "Direct Revised BER")
    _groww_ber_number(revised_regular, "Regular Revised BER")
    if direct is not None:
        plans["Direct"] = direct
    if regular is not None:
        plans["Regular"] = regular
    if not plans:
        raise ValueError("Groww BER notice has no numeric Current BER for Smallcap")
    return as_of.isoformat(), plans, int(notice.group(1)), signed.isoformat(), effective.isoformat()


def parse_groww_ber_pdf(content, today=None):
    return parse_groww_ber_text(_groww_pdf_text(content), today)


def _groww_disclosure(today=None):
    today = today or date.today()
    page, _, _ = fetch(GROWW_BER_PAGE, archive=False, max_bytes=2 * 1024 * 1024)
    candidates = _groww_ber_links(page, today)
    relevant = []
    for notice_number, source in candidates[:8]:
        raw, _, _ = fetch(source, archive=False, max_bytes=2 * 1024 * 1024)
        text = _groww_pdf_text(raw)
        if not re.search(r"Groww\s+Smallcap\s+Fund", text, re.I):
            continue
        day, plans, parsed_notice, signed, effective = parse_groww_ber_text(text, today)
        if parsed_notice != notice_number:
            raise ValueError("Groww BER notice title number and PDF identity do not match")
        relevant.append((day, notice_number, source, plans, signed, effective))

    if not relevant:
        raise ValueError("Groww BER notices contain no Smallcap Current BER observation")
    latest_day = max(item[0] for item in relevant)
    latest = [item for item in relevant if item[0] == latest_day]
    signatures = {(tuple(sorted(item[3].items())), item[4], item[5]) for item in latest}
    if len(signatures) != 1:
        raise ValueError(f"Groww BER notices conflict for latest observation {latest_day}")
    selected = max(latest, key=lambda item: item[1])
    _, notice_number, source, _, _, _ = selected

    content, content_hash, _ = fetch(source, max_bytes=2 * 1024 * 1024)
    day, plans, parsed_notice, signed, effective = parse_groww_ber_pdf(content, today)
    if parsed_notice != notice_number or day != latest_day:
        raise ValueError("Groww BER notice changed between discovery and archive fetch")
    return source, day, plans, content_hash, signed, effective


def groww(progress=lambda _: None, today=None):
    """Collect Groww Small Cap's explicitly published Current BER observations."""
    today = today or date.today()
    if not db.one("SELECT code FROM schemes WHERE family=? LIMIT 1", (GROWW_FAMILY,)):
        return "Groww Small Cap is not in the active universe"

    progress("Groww Small Cap expense ratios · official Current BER notice")
    source, day, plans, content_hash, signed, effective = _groww_disclosure(today)
    for plan, value in plans.items():
        db.metric(
            GROWW_FAMILY,
            plan,
            "base_expense_ratio",
            day,
            value,
            "% p.a. · reported by AMC",
            source,
            content_hash,
        )
    detail = ", ".join(f"{plan} {value:.2f}%" for plan, value in sorted(plans.items()))
    return (
        f"{GROWW_FAMILY}: official Current BER as of {day} ({detail}); "
        f"notice {signed}, revised BER effective {effective} not promoted"
    )


def _hsbc_key(secret):
    # CAMS' public browser bundle hashes these presentation-layer constants and
    # uses the first 32 hex characters as the UTF-8 AES-256 key bytes.
    return hashlib.sha256(secret.encode("utf-8")).hexdigest()[:32].encode("ascii")


def _hsbc_aes(data, secret, *, decrypt=False):
    openssl = shutil.which("openssl")
    if not openssl:
        raise ValueError("OpenSSL is required to decode the public HSBC TER download")
    command = [
        openssl,
        "enc",
        "-aes-256-cbc",
        "-K",
        _hsbc_key(secret).hex(),
        "-iv",
        _HSBC_IV.hex(),
        "-nosalt",
    ]
    if decrypt:
        command.append("-d")
    result = subprocess.run(
        command,
        input=data,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        timeout=15,
        check=False,
    )
    if result.returncode:
        raise ValueError("CAMS public HSBC TER transport could not be decoded")
    return result.stdout


def _hsbc_request_payload():
    payload = {
        "flag": "GET_UPD_MB_RESULT",
        "jobday": "hsbc_ter_report.xlsx",
        "browser": "Chrome",
        "device_id": "153.0.0.0",
        "os_id": "10",
        "application": "CAMSONLINE",
        "sub_application": "DIGITALADMIN",
        "deviceid": "desktop",
        "page_name": "/dnlresult/hsbc_ter_report.xlsx",
    }
    raw = json.dumps(payload, separators=(",", ":")).encode("utf-8")
    ciphertext = _hsbc_aes(raw, _HSBC_REQUEST_SECRET)
    return base64.b64encode(ciphertext).decode("ascii").replace("+", "-").replace("/", "_")


def _hsbc_decode_response(raw):
    if len(raw) > 16 * 1024 * 1024:
        raise ValueError("CAMS HSBC TER response is unexpectedly large")
    try:
        outer = json.loads(raw)
        if not isinstance(outer, str) or not outer:
            raise ValueError
        ciphertext = base64.b64decode(
            outer.replace("-", "+").replace("_", "/"), validate=True
        )
    except (json.JSONDecodeError, ValueError, TypeError) as exc:
        raise ValueError("CAMS HSBC TER response has an invalid encrypted envelope") from exc
    plain = _hsbc_aes(ciphertext, _HSBC_RESPONSE_SECRET, decrypt=True)
    try:
        payload = json.loads(plain.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError("CAMS HSBC TER response decrypted to invalid JSON") from exc
    status = payload.get("status")
    details = payload.get("detail")
    if (
        not isinstance(status, dict)
        or bool(status.get("errorflag"))
        or not isinstance(details, list)
        or len(details) != 1
        or not isinstance(details[0], dict)
    ):
        raise ValueError("CAMS HSBC TER service did not return one successful workbook")
    result = details[0].get("RESULT")
    if not isinstance(result, str) or not result:
        raise ValueError("CAMS HSBC TER response did not include workbook bytes")
    try:
        workbook = base64.b64decode(result, validate=True)
    except ValueError as exc:
        raise ValueError("CAMS HSBC TER workbook payload is invalid base64") from exc
    if not workbook.startswith(b"PK") or not 10_000 < len(workbook) <= 12 * 1024 * 1024:
        raise ValueError("CAMS HSBC TER payload is not a plausible XLSX workbook")
    return workbook


def _hsbc_workbook():
    public_url(HSBC_CAMS_API)
    encoded = _hsbc_request_payload()
    headers = {
        "User-Agent": "SmallcapLedger/1.0 (public AMC disclosure collection)",
        "Accept": "application/json, text/plain, */*",
        "Content-Type": "application/json",
        "Origin": "https://digital.camsonline.com",
        "Referer": HSBC_TER_LINK,
    }
    try:
        with httpx.Client(timeout=60, headers=headers, follow_redirects=False) as client:
            response = client.post(HSBC_CAMS_API, json={"data": encoded})
            response.raise_for_status()
            raw = response.content
    except httpx.HTTPError as exc:
        raise ValueError("CAMS HSBC TER service is unavailable") from exc
    return _hsbc_decode_response(raw)


def _hsbc_day(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def _hsbc_number(value, label):
    if value is None or str(value).strip() in ("", "-", "NA", "N/A"):
        raise ValueError(f"HSBC TER workbook is missing {label}")
    parsed = number(value)
    if not 0 <= parsed <= 5:
        raise ValueError(f"HSBC TER workbook {label} is outside the accepted range")
    return parsed


def parse_hsbc_workbook(content, today=None):
    """Return HSBC Small Cap's newest explicit BER and Total TER plan pair."""
    today = today or date.today()
    try:
        book = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError("HSBC TER disclosure is not a readable XLSX workbook") from exc
    try:
        if "TER" not in book.sheetnames:
            raise ValueError("HSBC TER workbook no longer contains the TER sheet")
        sheet = book["TER"]
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)
        title = next(rows, None)
        header = next(rows, None)
        if not title or str(title[0] or "").strip() != "Total Expense Ratio (TER) for HSBC Mutual Fund":
            raise ValueError("HSBC TER workbook title changed")
        if not header or tuple(str(x or "").strip() for x in header[:14]) != _HSBC_HEADER:
            raise ValueError("HSBC TER workbook columns changed")

        matches = defaultdict(list)
        for row in rows:
            if len(row) < 14:
                continue
            if str(row[0] or "").strip() != HSBC_SCHEME_CODE:
                continue
            if str(row[1] or "").strip() != HSBC_NSDL_CODE:
                continue
            if str(row[2] or "").strip() != HSBC_FAMILY:
                continue
            day = _hsbc_day(row[3])
            if day is None or day > today:
                continue
            matches[day.isoformat()].append(row)

        if not matches:
            raise ValueError("HSBC TER workbook contains no dated Small Cap rows")
        day = max(matches)
        if len(matches[day]) != 1:
            raise ValueError(f"HSBC TER workbook has duplicate Small Cap rows for {day}")
        row = matches[day][0]

        regular = {
            "base_expense_ratio": _hsbc_number(row[4], "Regular BER"),
            "brokerage": _hsbc_number(row[5], "Regular brokerage"),
            "transaction_cost": _hsbc_number(row[6], "Regular transaction cost"),
            "statutory_levies": _hsbc_number(row[7], "Regular statutory levies"),
            "ter": _hsbc_number(row[8], "Regular Total TER"),
        }
        direct = {
            "base_expense_ratio": _hsbc_number(row[9], "Direct BER"),
            "brokerage": _hsbc_number(row[10], "Direct brokerage"),
            "transaction_cost": _hsbc_number(row[11], "Direct transaction cost"),
            "statutory_levies": _hsbc_number(row[12], "Direct statutory levies"),
            "ter": _hsbc_number(row[13], "Direct Total TER"),
        }
        for plan, values in (("Regular", regular), ("Direct", direct)):
            if values["ter"] + 1e-9 < values["base_expense_ratio"]:
                raise ValueError(f"HSBC {plan} Total TER is below BER on {day}")
            component_total = (
                values["base_expense_ratio"]
                + values["brokerage"]
                + values["transaction_cost"]
                + values["statutory_levies"]
            )
            if abs(component_total - values["ter"]) > 0.02:
                raise ValueError(f"HSBC {plan} TER components do not reconcile on {day}")
        return day, {"Regular": regular, "Direct": direct}
    finally:
        book.close()


def hsbc(progress=lambda _: None, today=None):
    """Collect HSBC Small Cap's detailed TER workbook linked by its AMC factsheet."""
    today = today or date.today()
    if not db.one("SELECT code FROM schemes WHERE family=? LIMIT 1", (HSBC_FAMILY,)):
        return "HSBC Small Cap is not in the active universe"

    progress("HSBC Small Cap expense ratios · AMC-linked detailed TER workbook")
    try:
        workbook = _hsbc_workbook()
        content_hash = db.archive(workbook, _HSBC_XLSX_MIME)
        with db.connect() as connection:
            connection.execute(
                "INSERT INTO fetches(url,fetched_at,status,hash) VALUES(?,?,?,?)",
                (HSBC_TER_LINK, db.now(), "ok", content_hash),
            )
        day, plans = parse_hsbc_workbook(workbook, today)
        for plan, values in plans.items():
            for metric in (
                "base_expense_ratio",
                "brokerage",
                "transaction_cost",
                "statutory_levies",
                "ter",
            ):
                db.metric(
                    HSBC_FAMILY,
                    plan,
                    metric,
                    day,
                    values[metric],
                    "% p.a. · reported by AMC via AMC-linked CAMS workbook",
                    HSBC_TER_LINK,
                    content_hash,
                )
    except Exception as exc:
        with db.connect() as connection:
            connection.execute(
                "INSERT INTO fetches(url,fetched_at,status,detail) VALUES(?,?,?,?)",
                (HSBC_TER_LINK, db.now(), "error", str(exc)[:400]),
            )
        raise

    return (
        f"{HSBC_FAMILY}: official detailed BER/TER as of {day} "
        f"(Direct {plans['Direct']['base_expense_ratio']:.2f}%/"
        f"{plans['Direct']['ter']:.2f}% BER/TER)"
    )



def _icici_headers():
    return {
        "User-Agent": "SmallcapLedger/1.0 (public AMC disclosure collection)",
        "Accept": "application/json, text/plain, */*",
        "Content-Type": "application/json",
        "Origin": "https://www.icicipruamc.com",
        "Referer": ICICI_TER_PAGE,
        "env": "api",
        "baggage": "",
        "requestAPIId": str(uuid.uuid4()),
    }


def _icici_api(url, *, body=None):
    """Call the same unauthenticated financial-disclosure API as ICICI's SPA."""
    public_url(url)
    try:
        with httpx.Client(
            timeout=httpx.Timeout(60, connect=15),
            headers=_icici_headers(),
            follow_redirects=False,
        ) as client:
            response = client.post(url, json=body) if body is not None else client.get(url)
            response.raise_for_status()
            payload = response.json()
    except (httpx.HTTPError, json.JSONDecodeError, ValueError) as exc:
        raise ValueError("ICICI financial disclosure API is unavailable or invalid") from exc
    if not isinstance(payload, dict) or not isinstance(payload.get("success"), dict):
        raise ValueError("ICICI financial disclosure API response changed format")
    if payload.get("error"):
        raise ValueError("ICICI financial disclosure API returned an error")
    return payload["success"].get("data")


def _icici_financial_year(today):
    start = today.year if today.month >= 4 else today.year - 1
    return f"{start}-{start + 1}"


def _icici_category_ids(categories, today=None):
    today = today or date.today()
    if not isinstance(categories, list):
        raise ValueError("ICICI financial categories response is not a list")
    top = [
        row for row in categories
        if isinstance(row, dict)
        and row.get("isEnabled") is True
        and str((row.get("title") or {}).get("code", "")).strip() == ICICI_TER_CATEGORY_CODE
        and str((row.get("title") or {}).get("text", "")).strip() == ICICI_TER_TITLE
        and str(row.get("internalName", "")).strip() == "total-expense-ratio"
    ]
    if len(top) != 1:
        raise ValueError("ICICI Total Expense Ratio category is not uniquely identified")
    parent = top[0]
    sub = [
        row for row in (parent.get("subCategory") or [])
        if isinstance(row, dict)
        and row.get("isEnabled") is True
        and str((row.get("title") or {}).get("code", "")).strip() == ICICI_TER_SUBCATEGORY_CODE
        and str((row.get("title") or {}).get("text", "")).strip() == ICICI_TER_TITLE
        and str(row.get("internalName", "")).strip() == "Total Expense Ratio"
    ]
    if len(sub) != 1:
        raise ValueError("ICICI Total Expense Ratio subcategory is not uniquely identified")
    child = sub[0]
    filters = {
        str((item.get("key") or {}).get("code", "")).strip():
        {str(value.get("code", "")).strip() for value in (item.get("value") or []) if isinstance(value, dict)}
        for item in (child.get("filter") or [])
        if isinstance(item, dict)
    }
    if ICICI_TER_SHOW not in filters.get("SHOW", set()):
        raise ValueError("ICICI TER Details filter is no longer published")
    financial_year = _icici_financial_year(today)
    if financial_year not in filters.get("FINANCIAL_YEAR", set()):
        raise ValueError("ICICI current financial year is absent from TER filters")
    return str(parent.get("id")), str(child.get("id")), financial_year


def _icici_month(value):
    token = str(value or "").strip()
    for fmt in ("%B%Y", "%b%Y"):
        try:
            parsed = datetime.strptime(token, fmt).date()
            return parsed.year, parsed.month
        except ValueError:
            pass
    return None


def _icici_select_file(files, parent_id, child_id, financial_year, today=None):
    today = today or date.today()
    if not isinstance(files, list):
        raise ValueError("ICICI TER file response is not a list")
    found = defaultdict(list)
    for row in files:
        if not isinstance(row, dict):
            continue
        title = row.get("title") or {}
        title_text = str(title.get("text", "") if isinstance(title, dict) else "").strip()
        match = _ICICI_TITLE.fullmatch(title_text)
        path = str(row.get("url", "")).strip()
        file_name = path.rsplit("/", 1)[-1]
        file_match = _ICICI_FILE.fullmatch(file_name)
        if not match or not file_match:
            continue
        title_month = _icici_month(match.group(1) + match.group(2))
        file_month = _icici_month(file_match.group(1) + file_match.group(2))
        if title_month is None or file_month is None or title_month != file_month:
            continue
        year, month = title_month
        if (year, month) > (today.year, today.month):
            continue
        if (
            row.get("isEnabled") is not True
            or str(row.get("category", "")).strip() != ICICI_TER_CATEGORY_CODE
            or str(row.get("categoryName", "")).strip() != ICICI_TER_TITLE
            or str(row.get("level1Id", "")).strip() != parent_id
            or str(row.get("level2Id", "")).strip() != child_id
            or ICICI_TER_SHOW not in (row.get("SHOW") or [])
            or financial_year not in (row.get("FINANCIAL_YEAR") or [])
            or not path.startswith("/financials-disclosures-files/Files/Total Expense Ratio/")
            or not path.lower().endswith(".xlsx")
        ):
            continue
        found[(year, month)].append(row)
    if not found:
        raise ValueError("ICICI TER API returned no current TER Details workbook")
    key = max(found)
    if len(found[key]) != 1:
        raise ValueError("ICICI TER API returned duplicate workbooks for the latest month")
    row = found[key][0]
    source = ICICI_FILE_BASE + quote(str(row["url"]), safe="/")
    public_url(source)
    return source


def _icici_inline_strings(content):
    """Read rows from ICICI's XLSX sheet XML.

    The current workbook has valid inline-string rows but malformed worksheet
    dimension metadata, so ordinary openpyxl iteration exposes only row 1.
    Parsing the standard OOXML sheet cells retains the source exactly.
    """
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise ValueError("ICICI TER disclosure is not a valid XLSX package") from exc
    try:
        sheets = [
            name for name in archive.namelist()
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)
        ]
        if len(sheets) != 1:
            raise ValueError("ICICI TER workbook must contain exactly one worksheet")
        raw = archive.read(sheets[0])
        if len(raw) > 12 * 1024 * 1024:
            raise ValueError("ICICI TER worksheet XML is unexpectedly large")
        root = ET.fromstring(raw)
    except (KeyError, ET.ParseError) as exc:
        raise ValueError("ICICI TER worksheet XML is invalid") from exc
    finally:
        archive.close()
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rows = []
    for row in root.findall(".//x:sheetData/x:row", ns):
        values = {}
        for cell in row.findall("x:c", ns):
            ref = str(cell.attrib.get("r", ""))
            column = re.match(r"[A-Z]+", ref)
            if not column:
                continue
            parts = [node.text or "" for node in cell.findall(".//x:t", ns)]
            if parts:
                value = "".join(parts)
            else:
                node = cell.find("x:v", ns)
                value = "" if node is None else (node.text or "")
            values[column.group(0)] = value
        rows.append(values)
    return rows


def _icici_percent(value, label):
    raw = str(value or "").strip()
    if not raw or raw.upper() in ("NA", "N/A", "-"):
        raise ValueError(f"ICICI TER workbook is missing {label}")
    parsed = number(raw)
    if not 0 <= parsed <= 5:
        raise ValueError(f"ICICI TER workbook {label} is outside the accepted range")
    return parsed


def parse_icici_workbook(content, today=None):
    """Return the newest exact Small Cap BER and published Total TER pair."""
    today = today or date.today()
    rows = _icici_inline_strings(content)
    if len(rows) < 4:
        raise ValueError("ICICI TER workbook contains no usable rows")
    if rows[0].get("A", "").strip() != "Total Expense Ratio (TER) for Mutual Fund Schemes":
        raise ValueError("ICICI TER workbook title changed")
    if rows[1].get("C", "").strip() != "Regular Plan" or rows[1].get("H", "").strip() != "Direct Plan":
        raise ValueError("ICICI TER plan headings changed")
    for column, expected in _ICICI_HEADER.items():
        if rows[2].get(column, "").strip() != expected:
            raise ValueError(f"ICICI TER column {column} changed")

    matches = defaultdict(list)
    for row in rows[3:]:
        if row.get("A", "").strip() != ICICI_FAMILY:
            continue
        try:
            day = datetime.strptime(row.get("B", "").strip(), "%d/%m/%Y").date()
        except ValueError:
            continue
        if day <= today:
            matches[day.isoformat()].append(row)
    if not matches:
        raise ValueError("ICICI TER workbook contains no dated Small Cap rows")
    day = max(matches)
    if len(matches[day]) != 1:
        raise ValueError(f"ICICI TER workbook has duplicate Small Cap rows for {day}")
    row = matches[day][0]

    regular = {
        "base_expense_ratio": _icici_percent(row.get("C"), "Regular BER"),
        "brokerage": _icici_percent(row.get("D"), "Regular brokerage"),
        "transaction_cost": _icici_percent(row.get("E"), "Regular transaction cost"),
        "statutory_levies": _icici_percent(row.get("F"), "Regular statutory levies"),
        "ter": _icici_percent(row.get("G"), "Regular Total TER"),
    }
    direct = {
        "base_expense_ratio": _icici_percent(row.get("H"), "Direct BER"),
        "brokerage": _icici_percent(row.get("I"), "Direct brokerage"),
        "transaction_cost": _icici_percent(row.get("J"), "Direct transaction cost"),
        "statutory_levies": _icici_percent(row.get("K"), "Direct statutory levies"),
        "ter": _icici_percent(row.get("L"), "Direct Total TER"),
    }
    for plan, values in (("Regular", regular), ("Direct", direct)):
        if values["ter"] + 1e-9 < values["base_expense_ratio"]:
            raise ValueError(f"ICICI {plan} Total TER is below BER on {day}")
        component_total = (
            values["base_expense_ratio"]
            + values["brokerage"]
            + values["transaction_cost"]
            + values["statutory_levies"]
        )
        if abs(component_total - values["ter"]) > 0.02:
            raise ValueError(f"ICICI {plan} TER components do not reconcile on {day}")
    return day, {"Regular": regular, "Direct": direct}


def _icici_disclosure(today=None):
    today = today or date.today()
    categories = _icici_api(ICICI_CATEGORIES_API + "?userType=Investor")
    parent_id, child_id, financial_year = _icici_category_ids(categories, today)
    payload = {
        "categoryId": child_id,
        "userType": "Investor",
        "fileType": "All",
        "page": "1",
        "size": "20",
        "filter": [
            {"SHOW": [ICICI_TER_SHOW]},
            {"FINANCIAL_YEAR": [financial_year]},
        ],
        "search": "",
    }
    data = _icici_api(ICICI_FILES_API, body=payload)
    if not isinstance(data, dict):
        raise ValueError("ICICI TER file-list response changed format")
    source = _icici_select_file(
        data.get("files"),
        parent_id,
        child_id,
        financial_year,
        today,
    )
    workbook, content_hash, _ = fetch(source, max_bytes=5 * 1024 * 1024)
    if not workbook.startswith(b"PK"):
        raise ValueError("ICICI TER file is not an XLSX workbook")
    return source, workbook, content_hash


def icici(progress=lambda _: None, today=None):
    """Collect ICICI Prudential Small Cap's explicit BER/Total TER workbook."""
    today = today or date.today()
    if not db.one("SELECT code FROM schemes WHERE family=? LIMIT 1", (ICICI_FAMILY,)):
        return "ICICI Prudential Small Cap is not in the active universe"
    progress("ICICI Prudential Small Cap expense ratios · official TER Details workbook")
    source, workbook, content_hash = _icici_disclosure(today)
    day, plans = parse_icici_workbook(workbook, today)
    for plan, values in plans.items():
        for metric in (
            "base_expense_ratio",
            "brokerage",
            "transaction_cost",
            "statutory_levies",
            "ter",
        ):
            db.metric(
                ICICI_FAMILY,
                plan,
                metric,
                day,
                values[metric],
                "% p.a. · reported by AMC",
                source,
                content_hash,
            )
    return (
        f"{ICICI_FAMILY}: official BER/TER as of {day} "
        f"(Direct {plans['Direct']['base_expense_ratio']:.2f}%/"
        f"{plans['Direct']['ter']:.2f}% BER/TER)"
    )


def _invesco_percent(value, label):
    raw = str(value or "").strip()
    if not raw or raw.upper() in ("NA", "N/A", "-"):
        raise ValueError(f"Invesco TER API is missing {label}")
    parsed = number(raw)
    if not 0 <= parsed <= 5:
        raise ValueError(f"Invesco TER API {label} is outside the accepted range")
    return parsed


def parse_invesco_records(records, today=None):
    """Return the newest exact Small Cap BER and published Total TER pair."""
    today = today or date.today()
    if not isinstance(records, list):
        raise ValueError("Invesco TER API response changed format")
    matches = defaultdict(list)
    for row in records:
        if not isinstance(row, dict):
            continue
        if str(row.get("Scheme Name", "")).strip() != INVESCO_FAMILY:
            continue
        if str(row.get("NSDL Scheme Code", "")).strip() != INVESCO_NSDL_CODE:
            continue
        raw_day = str(row.get("TER Date(DD/MM/YYYY)", "")).strip()
        try:
            day = datetime.strptime(raw_day, "%d/%m/%Y").date()
        except ValueError:
            continue
        if day <= today:
            matches[day.isoformat()].append(row)
    if not matches:
        raise ValueError("Invesco TER API contains no dated Small Cap rows")
    day = max(matches)
    if len(matches[day]) != 1:
        raise ValueError(f"Invesco TER API has duplicate Small Cap rows for {day}")
    row = matches[day][0]
    plans = {}
    for plan, fields in _INVESCO_FIELDS.items():
        values = {
            metric: _invesco_percent(row.get(field), f"{plan} {metric}")
            for metric, field in fields.items()
        }
        if values["ter"] + 1e-9 < values["base_expense_ratio"]:
            raise ValueError(f"Invesco {plan} Total TER is below BER on {day}")
        component_total = (
            values["base_expense_ratio"]
            + values["brokerage"]
            + values["transaction_cost"]
            + values["statutory_levies"]
        )
        if abs(component_total - values["ter"]) > 0.02:
            raise ValueError(f"Invesco {plan} TER components do not reconcile on {day}")
        plans[plan] = values
    return day, plans


def _invesco_financial_year_start(day):
    return day.year if day.month >= 4 else day.year - 1


def _invesco_periods(today):
    first = today.replace(day=1)
    previous = (first - timedelta(days=1)).replace(day=1)
    return (first, previous)


def _invesco_json(raw, label):
    try:
        payload = json.loads(raw)
    except (TypeError, json.JSONDecodeError) as exc:
        raise ValueError(f"Invesco {label} API returned invalid JSON") from exc
    return payload


def _invesco_disclosure(today=None):
    today = today or date.today()
    plans_raw, _, _ = fetch(INVESCO_PLANS_API, archive=False, max_bytes=1024 * 1024)
    plans = _invesco_json(plans_raw, "plans")
    if not isinstance(plans, list) or plans.count(INVESCO_FAMILY) != 1:
        raise ValueError("Invesco Small Cap is not uniquely present in the public TER selector")

    selected_url = None
    for period in _invesco_periods(today):
        query = urlencode({
            "title": INVESCO_FAMILY,
            "fincialYear": _invesco_financial_year_start(period),
            "month": period.month,
        })
        url = INVESCO_TER_API + "?" + query
        raw, _, _ = fetch(url, archive=False, max_bytes=2 * 1024 * 1024)
        records = _invesco_json(raw, "TER")
        if not isinstance(records, list):
            raise ValueError("Invesco TER API response changed format")
        if not records:
            continue
        # A populated month for the exact title must contain the exact identity.
        parse_invesco_records(records, today)
        selected_url = url
        break
    if selected_url is None:
        raise ValueError("Invesco TER API returned no current or previous-month Small Cap rows")

    raw, content_hash, _ = fetch(selected_url, max_bytes=2 * 1024 * 1024)
    records = _invesco_json(raw, "TER")
    day, plans = parse_invesco_records(records, today)
    return selected_url, day, plans, content_hash


def invesco(progress=lambda _: None, today=None):
    """Collect Invesco India Small Cap's explicit BER and Total TER JSON rows."""
    today = today or date.today()
    if not db.one("SELECT code FROM schemes WHERE family=? LIMIT 1", (INVESCO_FAMILY,)):
        return "Invesco India Small Cap is not in the active universe"
    progress("Invesco India Small Cap expense ratios · official TER disclosure API")
    source, day, plans, content_hash = _invesco_disclosure(today)
    for plan, values in plans.items():
        for metric in (
            "base_expense_ratio",
            "brokerage",
            "transaction_cost",
            "statutory_levies",
            "ter",
        ):
            db.metric(
                INVESCO_FAMILY,
                plan,
                metric,
                day,
                values[metric],
                "% p.a. · reported by AMC",
                source,
                content_hash,
            )
    return (
        f"{INVESCO_FAMILY}: official BER/TER as of {day} "
        f"(Direct {plans['Direct']['base_expense_ratio']:.2f}%/"
        f"{plans['Direct']['ter']:.2f}% BER/TER)"
    )


def _jm_percent(value, label):
    if value is None or str(value).strip() in ("", "-", "NA", "N/A"):
        raise ValueError(f"JM TER API is missing {label}")
    parsed = number(value)
    if not 0 <= parsed <= 5:
        raise ValueError(f"JM TER API {label} is outside the accepted range")
    return parsed


def parse_jm_ter_records(records, today=None):
    """Return the newest exact JM Small Cap BER and published Total TER pair."""
    today = today or date.today()
    if not isinstance(records, list):
        raise ValueError("JM TER API response changed format")

    matches = defaultdict(list)
    for row in records:
        if not isinstance(row, dict):
            continue
        if str(row.get("Scheme", "")).strip() != JM_PUBLISHED_NAME:
            continue
        if str(row.get("Schemecode", "")).strip().upper() != JM_SCHEME_CODE:
            continue
        if str(row.get("NsdlSchemeCode", "")).strip() != JM_NSDL_CODE:
            continue
        raw_day = str(row.get("TERDate", "")).strip()
        try:
            day = datetime.fromisoformat(raw_day.replace("Z", "+00:00")).date()
        except ValueError:
            continue
        if day <= today:
            matches[day.isoformat()].append(row)

    if not matches:
        raise ValueError("JM TER API contains no dated Small Cap rows")
    day = max(matches)
    if len(matches[day]) != 1:
        raise ValueError(f"JM TER API has duplicate Small Cap rows for {day}")
    row = matches[day][0]

    plans = {}
    for plan, fields in _JM_FIELDS.items():
        values = {
            metric: _jm_percent(row.get(field), f"{plan} {metric}")
            for metric, field in fields.items()
        }
        if values["ter"] + 1e-9 < values["base_expense_ratio"]:
            raise ValueError(f"JM {plan} Total TER is below BER on {day}")
        component_total = (
            values["base_expense_ratio"]
            + values["brokerage"]
            + values["transaction_cost"]
            + values["statutory_levies"]
        )
        if abs(component_total - values["ter"]) > 0.02:
            raise ValueError(f"JM {plan} TER components do not reconcile on {day}")
        plans[plan] = values
    return day, plans


def _jm_ter_disclosure(today=None):
    today = today or date.today()
    raw, _, _ = fetch(
        JM_TER_API,
        body=JM_TER_REQUEST,
        archive=False,
        max_bytes=4 * 1024 * 1024,
    )
    records = jm_portfolios._decrypt(raw)
    day, plans = parse_jm_ter_records(records, today)

    # Archive the decoded first-party financial payload rather than the
    # AES-wrapped transport envelope. This is the exact JSON the public browser
    # renders after applying JM's published client-side transport key.
    evidence = json.dumps(
        records,
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=False,
    ).encode("utf-8")
    content_hash = db.archive(evidence, "application/json")
    with db.connect() as connection:
        connection.execute(
            "INSERT INTO fetches(url,fetched_at,status,hash) VALUES(?,?,?,?)",
            (JM_TER_API, db.now(), "ok", content_hash),
        )
    return day, plans, content_hash


def jm(progress=lambda _: None, today=None):
    """Collect JM Small Cap's explicit BER and Total TER from its public table API."""
    today = today or date.today()
    if not db.one("SELECT code FROM schemes WHERE family=? LIMIT 1", (JM_FAMILY,)):
        return "JM Small Cap is not in the active universe"

    progress("JM Small Cap expense ratios · official Scheme Expense Ratio API")
    try:
        day, plans, content_hash = _jm_ter_disclosure(today)
        for plan, values in plans.items():
            for metric in (
                "base_expense_ratio",
                "brokerage",
                "transaction_cost",
                "statutory_levies",
                "ter",
            ):
                db.metric(
                    JM_FAMILY,
                    plan,
                    metric,
                    day,
                    values[metric],
                    "% p.a. · reported by AMC",
                    JM_TER_API,
                    content_hash,
                )
    except Exception as exc:
        with db.connect() as connection:
            connection.execute(
                "INSERT INTO fetches(url,fetched_at,status,detail) VALUES(?,?,?,?)",
                (JM_TER_API, db.now(), "error", str(exc)[:400]),
            )
        raise

    return (
        f"{JM_FAMILY}: official BER/TER as of {day} "
        f"(Direct {plans['Direct']['base_expense_ratio']:.2f}%/"
        f"{plans['Direct']['ter']:.2f}% BER/TER)"
    )


def _mirae_header(value):
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def _mirae_percent(value, label):
    if value is None or str(value).strip() in ("", "-", "NA", "N/A"):
        raise ValueError(f"Mirae TER workbook is missing {label}")
    try:
        raw = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Mirae TER workbook has invalid {label}") from exc
    # The AMC's legacy XLS stores percentage-formatted cells as decimal fractions
    # (for example 0.0064 is displayed as 0.64%).
    if not 0 <= raw <= 0.05:
        raise ValueError(f"Mirae TER workbook {label} is outside the accepted range")
    return round(raw * 100, 8)


def parse_mirae_rows(rows, datemode=0, today=None):
    """Return the newest exact Mirae Small Cap BER and published Total TER pair."""
    today = today or date.today()
    matches = defaultdict(list)
    for row in rows:
        if len(row) < 13:
            continue
        if str(row[0] or "").strip() != MIRAE_NSDL_CODE:
            continue
        if str(row[1] or "").strip() != MIRAE_FAMILY:
            continue
        try:
            if isinstance(row[2], (int, float)):
                day = xlrd.xldate_as_datetime(row[2], datemode).date()
            else:
                day = datetime.strptime(str(row[2]).strip(), "%d/%m/%Y").date()
        except (ValueError, TypeError, xlrd.XLDateError):
            continue
        if day <= today:
            matches[day.isoformat()].append(row)

    if not matches:
        raise ValueError("Mirae TER workbook contains no dated Small Cap rows")
    day = max(matches)
    if len(matches[day]) != 1:
        raise ValueError(f"Mirae TER workbook has duplicate Small Cap rows for {day}")
    row = matches[day][0]

    regular = {
        "base_expense_ratio": _mirae_percent(row[3], "Regular BER"),
        "brokerage": _mirae_percent(row[4], "Regular brokerage"),
        "transaction_cost": _mirae_percent(row[5], "Regular transaction cost"),
        "statutory_levies": _mirae_percent(row[6], "Regular statutory levies"),
        "ter": _mirae_percent(row[7], "Regular Total TER"),
    }
    direct = {
        "base_expense_ratio": _mirae_percent(row[8], "Direct BER"),
        "brokerage": _mirae_percent(row[9], "Direct brokerage"),
        "transaction_cost": _mirae_percent(row[10], "Direct transaction cost"),
        "statutory_levies": _mirae_percent(row[11], "Direct statutory levies"),
        "ter": _mirae_percent(row[12], "Direct Total TER"),
    }
    for plan, values in (("Regular", regular), ("Direct", direct)):
        if values["ter"] + 1e-9 < values["base_expense_ratio"]:
            raise ValueError(f"Mirae {plan} Total TER is below BER on {day}")
        component_total = (
            values["base_expense_ratio"]
            + values["brokerage"]
            + values["transaction_cost"]
            + values["statutory_levies"]
        )
        if abs(component_total - values["ter"]) > 0.02:
            raise ValueError(f"Mirae {plan} TER components do not reconcile on {day}")
    return day, {"Regular": regular, "Direct": direct}


def parse_mirae_workbook(content, today=None):
    """Parse Mirae's explicit daily Total Expense Ratio legacy XLS workbook."""
    today = today or date.today()
    try:
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:
        raise ValueError("Mirae TER disclosure is not a readable XLS workbook") from exc
    try:
        if book.sheet_names() != ["Report"]:
            raise ValueError("Mirae TER workbook sheet layout changed")
        sheet = book.sheet_by_name("Report")
        if sheet.ncols != 13 or sheet.nrows < 3:
            raise ValueError("Mirae TER workbook dimensions changed")
        header1 = tuple(_mirae_header(sheet.cell_value(0, col)) for col in range(13))
        header2 = tuple(_mirae_header(sheet.cell_value(1, col)) for col in range(13))
        if header1 != _MIRAE_HEADER_1:
            raise ValueError("Mirae TER workbook top header changed")
        if header2 != _MIRAE_HEADER_2:
            raise ValueError("Mirae TER workbook metric columns changed")
        rows = [sheet.row_values(i, 0, 13) for i in range(2, sheet.nrows)]
        return parse_mirae_rows(rows, book.datemode, today)
    finally:
        book.release_resources()


def _mirae_dotnet_date(raw):
    match = re.fullmatch(r"/Date\((\d+)\)/", str(raw or "").strip())
    if not match:
        raise ValueError("Mirae TER metadata has an invalid PublishDate")
    return datetime.utcfromtimestamp(int(match.group(1)) / 1000).date()


def _mirae_select_download(payload, today=None):
    today = today or date.today()
    if not isinstance(payload, dict) or payload.get("ReturnCode") != "0":
        raise ValueError("Mirae TER download API returned an unsuccessful response")
    records = payload.get("Data")
    if not isinstance(records, list):
        raise ValueError("Mirae TER download API response changed format")

    matches = defaultdict(list)
    for item in records:
        if not isinstance(item, dict):
            continue
        title_match = _MIRAE_TITLE.fullmatch(str(item.get("Title") or "").strip())
        path = str(item.get("URL") or "").strip()
        file_match = _MIRAE_FILE.fullmatch(path.rsplit("/", 1)[-1])
        if not title_match or not file_match:
            continue
        if not path.startswith("/DailyUploads/TotalExpenseRatio/"):
            continue
        title_day = datetime.strptime(title_match.group(1), "%d %b %Y").date()
        file_day = datetime.strptime(file_match.group(1), "%d%m%Y").date()
        publish_day = _mirae_dotnet_date(item.get("PublishDate"))
        if title_day != file_day or publish_day != file_day or file_day > today:
            continue
        source = MIRAE_TER_BASE + path
        public_url(source)
        matches[file_day.isoformat()].append(source)

    if not matches:
        raise ValueError("Mirae TER download API contains no current valid workbook")
    day = max(matches)
    if len(matches[day]) != 1:
        raise ValueError(f"Mirae TER download API has duplicate workbooks for {day}")
    return day, matches[day][0]


def _mirae_disclosure(today=None):
    today = today or date.today()
    start = today - timedelta(days=35)
    request = {
        "request": {
            "modulename": "TotalExpenseRatio",
            "title": "",
            "fromdate": start.isoformat(),
            "todate": today.isoformat(),
            "pgno": 1,
            "pgsize": 100,
        }
    }
    raw, _, _ = fetch(
        MIRAE_TER_API,
        body=request,
        archive=False,
        max_bytes=4 * 1024 * 1024,
    )
    try:
        payload = json.loads(raw)
    except (TypeError, json.JSONDecodeError) as exc:
        raise ValueError("Mirae TER download API returned invalid JSON") from exc
    metadata_day, source = _mirae_select_download(payload, today)
    workbook, content_hash, _ = fetch(source, max_bytes=5 * 1024 * 1024)
    if not workbook.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        raise ValueError("Mirae TER source is not a legacy XLS workbook")
    day, plans = parse_mirae_workbook(workbook, today)
    if day != metadata_day:
        raise ValueError("Mirae TER workbook latest date does not match its publication metadata")
    return source, day, plans, content_hash


def mirae(progress=lambda _: None, today=None):
    """Collect Mirae Asset Small Cap's explicit BER and Total TER workbook."""
    today = today or date.today()
    if not db.one("SELECT code FROM schemes WHERE family=? LIMIT 1", (MIRAE_FAMILY,)):
        return "Mirae Asset Small Cap is not in the active universe"

    progress("Mirae Asset Small Cap expense ratios · official daily Total Expense Ratio workbook")
    source, day, plans, content_hash = _mirae_disclosure(today)
    for plan, values in plans.items():
        for metric in (
            "base_expense_ratio",
            "brokerage",
            "transaction_cost",
            "statutory_levies",
            "ter",
        ):
            db.metric(
                MIRAE_FAMILY,
                plan,
                metric,
                day,
                values[metric],
                "% p.a. · reported by AMC",
                source,
                content_hash,
            )
    return (
        f"{MIRAE_FAMILY}: official BER/TER as of {day} "
        f"(Direct {plans['Direct']['base_expense_ratio']:.2f}%/"
        f"{plans['Direct']['ter']:.2f}% BER/TER)"
    )


def _mahindra_decrypt_downloads(raw):
    """Decode the public Downloads API exactly as Mahindra's browser client does."""
    try:
        outer = json.loads(raw)
        payload = outer.get("payload")
        if not isinstance(payload, str) or not payload:
            raise ValueError
        ciphertext = base64.b64decode(payload, validate=True)
    except (json.JSONDecodeError, ValueError, TypeError) as exc:
        raise ValueError("Mahindra downloads API returned an invalid encrypted envelope") from exc

    if len(ciphertext) > 4 * 1024 * 1024:
        raise ValueError("Mahindra downloads API response is unexpectedly large")
    openssl = shutil.which("openssl")
    if not openssl:
        raise ValueError("OpenSSL is required to decode Mahindra public browser API data")
    result = subprocess.run(
        [
            openssl,
            "enc",
            "-d",
            "-aes-256-cbc",
            "-K",
            _MAHINDRA_AES_KEY.hex(),
            "-iv",
            _MAHINDRA_AES_IV.hex(),
        ],
        input=ciphertext,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        timeout=15,
        check=False,
    )
    if result.returncode:
        raise ValueError("Mahindra public downloads payload could not be decoded")
    try:
        decoded = json.loads(result.stdout.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError("Mahindra public downloads payload decrypted to invalid JSON") from exc
    if (
        not isinstance(decoded, dict)
        or decoded.get("status") != 1
        or not isinstance(decoded.get("data"), list)
    ):
        raise ValueError("Mahindra downloads API decoded response changed format")
    return decoded["data"]


def _mahindra_financial_year_title(today):
    start = today.year if today.month >= 4 else today.year - 1
    return f"TOTAL EXPENSE RATIO - {start}-{str(start + 1)[-2:]}"


def _mahindra_exact_children(nodes, name):
    return [
        node
        for node in (nodes or [])
        if isinstance(node, dict)
        and str(node.get("categoryName") or "").strip() == name
    ]


def _mahindra_select_ter_file(tree, today=None):
    """Select the exact current-financial-year TER workbook from AMC metadata."""
    today = today or date.today()
    top = _mahindra_exact_children(tree, _MAHINDRA_TOP_CATEGORY)
    if len(top) != 1:
        raise ValueError("Mahindra Mandatory Disclosures category is not uniquely identified")
    ter_group = _mahindra_exact_children(
        top[0].get("subcategories"), _MAHINDRA_TER_CATEGORY
    )
    if len(ter_group) != 1:
        raise ValueError("Mahindra TER disclosure group is not uniquely identified")
    ter = _mahindra_exact_children(
        ter_group[0].get("subcategories"), _MAHINDRA_TER_SUBCATEGORY
    )
    if len(ter) != 1:
        raise ValueError("Mahindra Total Expense Ratio category is not uniquely identified")

    wanted = _mahindra_financial_year_title(today)
    matches = []
    for row in ter[0].get("files") or []:
        if not isinstance(row, dict):
            continue
        if str(row.get("title") or "").strip().upper() != wanted:
            continue
        url = str(row.get("fileUrl") or "").strip()
        parsed = urlparse(url)
        if (
            parsed.scheme != "https"
            or parsed.hostname not in ("www.mahindramanulife.com", "cms.mahindramanulife.com")
            or not parsed.path.lower().endswith(".xlsx")
            or "/uploads/download/" not in parsed.path.lower()
        ):
            continue
        public_url(url)
        matches.append(url)
    if len(matches) != 1:
        raise ValueError("Mahindra current financial-year TER workbook is not uniquely identified")
    return matches[0]


def _mahindra_percent(value, label):
    if value is None or str(value).strip() in ("", "-", "NA", "N/A"):
        raise ValueError(f"Mahindra TER workbook is missing {label}")
    parsed = number(value)
    if not 0 <= parsed <= 5:
        raise ValueError(f"Mahindra TER workbook {label} is outside the accepted range")
    return parsed


def parse_mahindra_workbook(content, today=None):
    """Return the newest exact Small Cap BER and AMC-published Total TER pair."""
    today = today or date.today()
    try:
        book = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError("Mahindra TER disclosure is not a readable XLSX workbook") from exc

    try:
        if book.sheetnames != ["Sheet1"]:
            raise ValueError("Mahindra TER workbook sheet layout changed")
        sheet = book["Sheet1"]
        rows = sheet.iter_rows(values_only=True)
        header1 = next(rows, None)
        header2 = next(rows, None)
        clean1 = tuple("" if v is None else str(v).strip() for v in (header1 or ())[:13])
        clean2 = tuple("" if v is None else str(v).strip() for v in (header2 or ())[:13])
        if clean1 != _MAHINDRA_HEADER_1:
            raise ValueError("Mahindra TER workbook top header changed")
        if clean2 != _MAHINDRA_HEADER_2:
            raise ValueError("Mahindra TER workbook metric columns changed")

        matches = defaultdict(list)
        for row in rows:
            if len(row) < 13:
                continue
            if str(row[0] or "").strip() != MAHINDRA_NSDL_CODE:
                continue
            if str(row[1] or "").strip() != MAHINDRA_FAMILY:
                continue
            try:
                day = datetime.strptime(str(row[2] or "").strip(), "%d-%b-%Y").date()
            except ValueError:
                continue
            if day <= today:
                matches[day.isoformat()].append(row)

        if not matches:
            raise ValueError("Mahindra TER workbook contains no dated Small Cap rows")
        day = max(matches)
        if len(matches[day]) != 1:
            raise ValueError(f"Mahindra TER workbook has duplicate Small Cap rows for {day}")
        row = matches[day][0]

        regular = {
            "base_expense_ratio": _mahindra_percent(row[3], "Regular BER"),
            "brokerage": _mahindra_percent(row[4], "Regular brokerage"),
            "transaction_cost": _mahindra_percent(row[5], "Regular transaction cost"),
            "statutory_levies": _mahindra_percent(row[6], "Regular statutory levies"),
            "ter": _mahindra_percent(row[7], "Regular Total TER"),
        }
        direct = {
            "base_expense_ratio": _mahindra_percent(row[8], "Direct BER"),
            "brokerage": _mahindra_percent(row[9], "Direct brokerage"),
            "transaction_cost": _mahindra_percent(row[10], "Direct transaction cost"),
            "statutory_levies": _mahindra_percent(row[11], "Direct statutory levies"),
            "ter": _mahindra_percent(row[12], "Direct Total TER"),
        }
        for plan, values in (("Regular", regular), ("Direct", direct)):
            if values["ter"] + 1e-9 < values["base_expense_ratio"]:
                raise ValueError(f"Mahindra {plan} Total TER is below BER on {day}")
            component_total = (
                values["base_expense_ratio"]
                + values["brokerage"]
                + values["transaction_cost"]
                + values["statutory_levies"]
            )
            if abs(component_total - values["ter"]) > 0.02:
                raise ValueError(f"Mahindra {plan} TER components do not reconcile on {day}")
        return day, {"Regular": regular, "Direct": direct}
    finally:
        book.close()


def _mahindra_disclosure(today=None):
    today = today or date.today()
    raw, _, _ = fetch(
        MAHINDRA_DOWNLOADS_API,
        archive=False,
        max_bytes=4 * 1024 * 1024,
    )
    tree = _mahindra_decrypt_downloads(raw)
    source = _mahindra_select_ter_file(tree, today)
    workbook, content_hash, _ = fetch(source, max_bytes=5 * 1024 * 1024)
    if not workbook.startswith(b"PK"):
        raise ValueError("Mahindra TER source is not an XLSX workbook")
    day, plans = parse_mahindra_workbook(workbook, today)
    return source, day, plans, content_hash


def mahindra(progress=lambda _: None, today=None):
    """Collect Mahindra Manulife Small Cap's explicit BER and Total TER workbook."""
    today = today or date.today()
    if not db.one("SELECT code FROM schemes WHERE family=? LIMIT 1", (MAHINDRA_FAMILY,)):
        return "Mahindra Manulife Small Cap is not in the active universe"

    progress("Mahindra Manulife Small Cap expense ratios · official TER workbook")
    source, day, plans, content_hash = _mahindra_disclosure(today)
    for plan, values in plans.items():
        for metric in (
            "base_expense_ratio",
            "brokerage",
            "transaction_cost",
            "statutory_levies",
            "ter",
        ):
            db.metric(
                MAHINDRA_FAMILY,
                plan,
                metric,
                day,
                values[metric],
                "% p.a. · reported by AMC",
                source,
                content_hash,
            )
    return (
        f"{MAHINDRA_FAMILY}: official BER/TER as of {day} "
        f"(Direct {plans['Direct']['base_expense_ratio']:.2f}%/"
        f"{plans['Direct']['ter']:.2f}% BER/TER)"
    )


def update(progress=lambda _: None):
    results = []
    errors = []
    for collector in (canara, groww, hsbc, icici, invesco, jm, mahindra, mirae):
        try:
            results.append(collector(progress))
        except Exception as exc:
            errors.append(f"{collector.__name__}: {str(exc)[:220]}")
    if errors:
        raise ValueError("; ".join(results + errors))
    return "; ".join(results)
